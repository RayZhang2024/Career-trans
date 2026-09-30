import io
import json
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest
from sqlalchemy import select

from app.api.deps import get_user_cv_ingestion_service
from app.main import app
from app.models.candidate_cv_ingestion import (
    CandidateCVIngestionDraft,
    CandidateCVReviewBaseline,
    CandidateEvidenceRecord,
    CandidateStructuredProfile,
)
from app.models.user import User
from app.models.candidate_adviser import CandidateAdviserClarificationRecord
from app.models.candidate_profile_revision import CandidateProfileRevisionRecord
from app.schemas.cv_ingestion import CandidateCVData
from app.schemas.candidate_adviser import ClarificationAnswerKind, ClarificationInterpretation
from app.schemas.candidate_adviser_profile_proposal import SkillProposalUpdate
from app.services.cv_file_extraction_service import CVFileExtractionService
from app.services.canonical_candidate_read_service import CanonicalCandidateReadService
from app.services.cv_ingestion_service import CVIngestionReadService, CVIngestionService
from app.services.cv_interpretation_service import SemanticCVInterpreter
from app.services.cv_merge_service import CVMergeService
from app.services.active_candidate_evidence import ActiveCandidateEvidenceResolver
from app.services.candidate_adviser_profile_proposal import CandidateAdviserProfileProposalService
from app.services.profile_revision_service import CandidateProfileRevisionService, ProfileRevisionStale


class FakeInterpreter:
    def __init__(self) -> None:
        self.calls = 0

    def interpret(self, documents):
        self.calls += 1
        return CandidateCVData.model_validate({"skills": [{"name": "Python"}]})


class InvalidResponseClient:
    class responses:
        @staticmethod
        def create(**_kwargs):
            return type("Response", (), {"output_text": "not valid JSON"})()


class CapturingResponseClient:
    def __init__(self, output: str) -> None:
        self.output = output
        self.kwargs: dict[str, object] | None = None

    class _Responses:
        def __init__(self, parent) -> None:
            self._parent = parent

        def create(self, **kwargs):
            self._parent.kwargs = kwargs
            return type("Response", (), {"output_text": self._parent.output})()

    @property
    def responses(self):
        return self._Responses(self)


class FixedCVInterpreter:
    def __init__(self, output: CandidateCVData) -> None:
        self.output = output

    def interpret(self, documents):
        self.documents = documents
        return self.output.model_copy(deep=True)


def _benchmark_cv_bytes() -> bytes:
    return (Path(__file__).parent / "fixtures" / "issue_254_benchmark_cv.md").read_bytes()


def _evidence_output(document_sha256: str, segment_ids: list[str], *, title: str = "Senior software engineer") -> CandidateCVData:
    return CandidateCVData.model_validate({
        "evidence": [{
            "evidence_type": "employment",
            "title": title,
            "text": "Built a document review service and reduced manual processing time by 40%.",
            "skills": ["Python"],
            "provenance": [{"document_sha256": document_sha256, "segment_ids": segment_ids}],
        }]
    })


def _pdf_with_pages(*pages: str) -> bytes:
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {
                        NameObject("/F1"): DictionaryObject(
                            {
                                NameObject("/Type"): NameObject("/Font"),
                                NameObject("/Subtype"): NameObject("/Type1"),
                                NameObject("/BaseFont"): NameObject("/Helvetica"),
                            }
                        )
                    }
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 100 700 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _auth(client, email: str) -> dict[str, str]:
    credentials = {"email": email, "password": "strong-password"}
    assert client.post("/api/v1/auth/register", json=credentials).status_code == 201
    return {"Authorization": f"Bearer {client.post('/api/v1/auth/login', json=credentials).json()['access_token']}"}


def _json_cv() -> dict:
    return {
        "employment": [{"employer": "Example", "title": "Engineer", "description": "Built reliable systems."}],
        "skills": [{"name": "Python", "category": "technical"}],
        "projects": [{"name": "Platform", "description": "Delivered a service.", "skills": ["Python"]}],
        "achievements": [{"text": "Improved reliability."}],
        "evidence": [{"evidence_type": "employment", "title": "Engineer at Example", "text": "Built reliable systems.", "skills": ["Python"]}],
    }


def test_json_upload_import_is_deterministic_reviewable_and_requires_confirmation(client, db_session) -> None:
    interpreter = FakeInterpreter()
    service = CVIngestionService(db_session, interpreter=interpreter)
    app.dependency_overrides[get_user_cv_ingestion_service] = lambda: service
    try:
        headers = _auth(client, "cv-json@example.com")
        upload = client.post("/api/v1/cv-ingestion/upload", headers=headers, files=[("files", ("cv.json", json.dumps(_json_cv()), "application/json"))])
        assert upload.status_code == 201
        draft_id = upload.json()["id"]
        assert db_session.scalars(select(CandidateEvidenceRecord)).all() == []
        review = client.post(f"/api/v1/cv-ingestion/{draft_id}/interpret", headers=headers)
        assert review.status_code == 200
        assert review.json()["state"] == "review_ready"
        assert interpreter.calls == 0
        assert db_session.scalars(select(CandidateEvidenceRecord)).all() == []
        assert db_session.scalars(select(CandidateStructuredProfile)).all() == []
        confirmed = client.post(f"/api/v1/cv-ingestion/{draft_id}/confirm", headers=headers)
        assert confirmed.status_code == 200
        # Semantic claim plus deterministic current employment evidence.
        assert confirmed.json()["confirmed_evidence_count"] == 2
        record = db_session.scalar(select(CandidateEvidenceRecord))
        assert record is not None and record.user_id
        assert db_session.scalar(select(CandidateStructuredProfile)) is not None
        repeated = client.post(f"/api/v1/cv-ingestion/{draft_id}/confirm", headers=headers)
        assert repeated.status_code == 200
        assert repeated.json()["confirmed_evidence_count"] == 0
        assert len(db_session.scalars(select(CandidateEvidenceRecord)).all()) == 2
    finally:
        app.dependency_overrides.pop(get_user_cv_ingestion_service, None)


def test_real_cv_confirmation_makes_transferred_profile_revision_stale(client, db_session) -> None:
    email = "cv-confirmation-stales-transferred-revision@example.com"
    headers = _auth(client, email)
    user = db_session.scalar(select(User).where(User.email == email))
    initial = CandidateCVData(skills=[{"name": "Current A"}])
    db_session.add(CandidateStructuredProfile(
        user_id=user.id,
        structured_json=json.dumps(initial.model_dump(mode="json"), sort_keys=True),
    ))
    ActiveCandidateEvidenceResolver(db_session).resolve(user.id, initial)
    clarification_id = "f" * 64
    interpretation = ClarificationInterpretation(
        answer_kind=ClarificationAnswerKind.CAREER_FACT,
        confirmed_context_summary="A confirmed career fact supports the proposed skill.",
        proposed_evidence=[],
    )
    db_session.add(CandidateAdviserClarificationRecord(
        user_id=user.id,
        clarification_id=clarification_id,
        question_key="e" * 64,
        origin_assessment_fingerprint="d" * 64,
        question_text="What skill did you use?",
        question_source_references_json="[]",
        priority_index=0,
        answer_text="Rust",
        interpretation_json=json.dumps(interpretation.model_dump(mode="json"), sort_keys=True),
        status="confirmed",
    ))
    db_session.commit()
    proposals = CandidateAdviserProfileProposalService(db_session)
    proposal = proposals.materialize_from_confirmed_clarification(
        user.id,
        clarification_id,
        SkillProposalUpdate(operation="add", section="skills", item={"name": "Adviser proposal"}),
    )
    transfer = proposals.transfer_to_profile_revision(
        user.id, proposal.id, expected_revision=proposal.revision
    )
    linked_revision_id = transfer.profile_revision.id
    original_base_fingerprint = db_session.get(
        CandidateProfileRevisionRecord, linked_revision_id
    ).base_structured_fingerprint

    ingestion = CVIngestionService(db_session, interpreter=FakeInterpreter())
    app.dependency_overrides[get_user_cv_ingestion_service] = lambda: ingestion
    newer_cv = _json_cv()
    newer_cv["skills"] = [{"name": "CV B current skill", "category": "technical"}]
    try:
        uploaded = client.post(
            "/api/v1/cv-ingestion/upload",
            headers=headers,
            files=[("files", ("newer-cv.json", json.dumps(newer_cv), "application/json"))],
        )
        assert uploaded.status_code == 201, uploaded.text
        draft_id = uploaded.json()["id"]
        reviewed_cv = client.post(f"/api/v1/cv-ingestion/{draft_id}/interpret", headers=headers)
        assert reviewed_cv.status_code == 200, reviewed_cv.text
        assert reviewed_cv.json()["state"] == "review_ready"
        confirmed_cv = client.post(f"/api/v1/cv-ingestion/{draft_id}/confirm", headers=headers)
        assert confirmed_cv.status_code == 200, confirmed_cv.text
        assert ingestion._interpreter.calls == 0
    finally:
        app.dependency_overrides.pop(get_user_cv_ingestion_service, None)

    current_row = db_session.scalar(select(CandidateStructuredProfile).where(
        CandidateStructuredProfile.user_id == user.id
    ))
    current = CandidateCVData.model_validate_json(current_row.structured_json)
    assert "CV B current skill" in [skill.name for skill in current.skills]
    stored_proposal = proposals.get_for_user(user.id, proposal.id)
    assert stored_proposal.state == "transferred"
    linked_revision = db_session.get(CandidateProfileRevisionRecord, linked_revision_id)
    assert linked_revision is not None and linked_revision.state == "draft"
    with pytest.raises(ProfileRevisionStale, match="structured"):
        CandidateProfileRevisionService(db_session).review(
            user.id, linked_revision_id, expected_revision=transfer.profile_revision.revision
        )
    db_session.refresh(linked_revision)
    assert linked_revision.state == "draft"
    assert linked_revision.revision == transfer.profile_revision.revision
    assert linked_revision.base_structured_fingerprint == original_base_fingerprint


def test_cv_drafts_are_scoped_to_authenticated_user(client, db_session) -> None:
    service = CVIngestionService(db_session, interpreter=FakeInterpreter())
    app.dependency_overrides[get_user_cv_ingestion_service] = lambda: service
    try:
        first = _auth(client, "first-cv@example.com")
        second = _auth(client, "second-cv@example.com")
        draft_id = client.post("/api/v1/cv-ingestion/upload", headers=first, files=[("files", ("cv.md", "# CV\nPython engineer", "text/markdown"))]).json()["id"]
        assert client.get(f"/api/v1/cv-ingestion/{draft_id}", headers=second).status_code == 404
    finally:
        app.dependency_overrides.pop(get_user_cv_ingestion_service, None)


def test_cv_history_is_bounded_user_scoped_and_metadata_only(client, db_session) -> None:
    first = _auth(client, "history-first@example.com")
    second = _auth(client, "history-second@example.com")
    for name in ("older.md", "newer.md"):
        response = client.post("/api/v1/cv-ingestion/upload", headers=first, files=[("files", (name, f"# {name}\nprivate source text", "text/markdown"))])
        assert response.status_code == 201
    other = client.post("/api/v1/cv-ingestion/upload", headers=second, files=[("files", ("other.md", "other", "text/markdown"))])
    assert other.status_code == 201

    assert client.get("/api/v1/cv-ingestion").status_code == 401
    response = client.get("/api/v1/cv-ingestion?limit=1", headers=first)
    assert response.status_code == 200
    payload = response.json()
    assert payload["limit"] == 1 and payload["truncated"] is True
    assert payload["items"][0]["filenames"] == ["newer.md"]
    assert payload["items"][0]["document_count"] == 1
    assert "segments" not in payload["items"][0]
    assert "private source text" not in response.text
    assert len(client.get("/api/v1/cv-ingestion", headers=second).json()["items"]) == 1
    assert client.get("/api/v1/cv-ingestion?limit=51", headers=first).status_code == 422


def test_cv_history_service_reads_without_sql_writes(db_session) -> None:
    user = User(email="read-only-history@example.com", password_hash="unused")
    db_session.add(user)
    db_session.commit()
    CVIngestionService(db_session).upload(user.id, [("cv.md", "text/markdown", b"source")])
    statements: list[str] = []
    from sqlalchemy import event

    def capture(_conn, _cursor, statement, *_args):
        statements.append(statement.lstrip().split(None, 1)[0].upper())

    event.listen(db_session.bind, "before_cursor_execute", capture)
    try:
        result = CVIngestionReadService(db_session).list(user.id)
    finally:
        event.remove(db_session.bind, "before_cursor_execute", capture)
    assert len(result.items) == 1
    assert statements and all(item not in {"INSERT", "UPDATE", "DELETE", "REPLACE"} for item in statements)


def test_confirmation_rolls_back_profile_evidence_and_draft_on_reconciliation_failure(db_session, monkeypatch) -> None:
    draft = CVIngestionService(db_session, interpreter=FakeInterpreter()).upload(
        "rollback-user", [("cv.json", "application/json", json.dumps(_json_cv()).encode())]
    )
    service = CVIngestionService(db_session, interpreter=FakeInterpreter())
    service.interpret("rollback-user", draft.id)

    def fail_resolution(_self, _user_id, _data):
        raise RuntimeError("forced evidence reconciliation failure")

    monkeypatch.setattr("app.services.cv_ingestion_service.ActiveCandidateEvidenceResolver.resolve", fail_resolution)
    with pytest.raises(RuntimeError, match="forced evidence"):
        service.confirm("rollback-user", draft.id)
    db_session.expire_all()
    assert db_session.scalar(select(CandidateStructuredProfile).where(CandidateStructuredProfile.user_id == "rollback-user")) is None
    assert db_session.scalars(select(CandidateEvidenceRecord).where(CandidateEvidenceRecord.user_id == "rollback-user")).all() == []
    persisted = db_session.scalar(select(CandidateCVIngestionDraft).where(CandidateCVIngestionDraft.id == draft.id))
    assert persisted is not None and persisted.state == "review_ready"


def test_confirmed_evidence_and_context_are_isolated_per_user(client, db_session) -> None:
    service = CVIngestionService(db_session, interpreter=FakeInterpreter())
    app.dependency_overrides[get_user_cv_ingestion_service] = lambda: service
    try:
        first_headers = _auth(client, "context-first@example.com")
        second_headers = _auth(client, "context-second@example.com")
        first_id = db_session.scalar(select(User.id).where(User.email == "context-first@example.com"))
        second_id = db_session.scalar(select(User.id).where(User.email == "context-second@example.com"))
        assert first_id and second_id
        draft = service.upload(first_id, [("cv.json", "application/json", json.dumps(_json_cv()).encode())])
        service.interpret(first_id, draft.id)
        service.confirm(first_id, draft.id)
        reader = CanonicalCandidateReadService(db_session)
        first_context = reader.candidate_context(reader.read(first_id), require_complete_evidence=True)
        second_context = reader.candidate_context(reader.read(second_id), require_complete_evidence=True)
        assert first_context is not None and first_context.evidence
        assert second_context is not None and second_context.evidence == []
        assert client.get(f"/api/v1/cv-ingestion/{draft.id}", headers=second_headers).status_code == 404
        assert first_headers
    finally:
        app.dependency_overrides.pop(get_user_cv_ingestion_service, None)


def test_markdown_uses_semantic_interpreter_and_context_uses_confirmed_data(db_session, runtime_snapshot_a) -> None:
    interpreter = FakeInterpreter()
    service = CVIngestionService(db_session, interpreter=interpreter, runtime_snapshot=runtime_snapshot_a)
    draft = service.upload("user-1", [("cv.md", "text/markdown", b"Python engineer")])
    service.interpret("user-1", draft.id)
    assert interpreter.calls == 1
    service.confirm("user-1", draft.id)
    reader = CanonicalCandidateReadService(db_session)
    context = reader.candidate_context(reader.read("user-1"), require_complete_evidence=True)
    assert context is not None
    assert context.skills_text == "Python"
    assert context.career_strategy_text == ""


def test_docx_extraction_and_conservative_merge_preserve_distinct_evidence() -> None:
    document = Document()
    document.add_heading("Experience", level=1)
    document.add_paragraph("Software engineer")
    buffer = io.BytesIO()
    document.save(buffer)
    extracted = CVFileExtractionService().extract(filename="cv.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", content=buffer.getvalue())
    assert extracted.segments[0].text == "Experience\nSoftware engineer"
    assert extracted.segments[0].heading == "Experience"
    merged = CVMergeService().merge([
        CandidateCVData.model_validate({"skills": [{"name": "Python"}], "evidence": [{"evidence_type": "project", "title": "A", "text": "Built A", "provenance": [{"document_sha256": "a", "segment_ids": ["a:1"]}]}]}),
        CandidateCVData.model_validate({"skills": [{"name": "python"}, {"name": "Rust"}], "evidence": [{"evidence_type": "project", "title": "A", "text": "Built A", "provenance": [{"document_sha256": "b", "segment_ids": ["b:1"]}]}, {"evidence_type": "project", "title": "B", "text": "Built B"}]}),
    ])
    assert [item.name for item in merged.skills] == ["Python", "Rust"]
    assert [item.title for item in merged.evidence] == ["A", "B"]
    assert {item.document_sha256 for item in merged.evidence[0].provenance} == {"a", "b"}
    assert {tuple(item.segment_ids) for item in merged.evidence[0].provenance} == {("a:1",), ("b:1",)}


def test_file_validation_rejects_unreadable_or_non_schema_input() -> None:
    extractor = CVFileExtractionService()
    try:
        extractor.extract(filename="cv.exe", content_type="application/octet-stream", content=b"not a CV")
    except ValueError as exc:
        assert "Unsupported" in str(exc)
    else:
        raise AssertionError("Unsupported files must be rejected")
    document = extractor.extract(filename="cv.json", content_type="application/json", content=b"[]")
    try:
        extractor.parsed_json(document)
    except ValueError as exc:
        assert "object" in str(exc)
    else:
        raise AssertionError("Non-object JSON must be rejected")


def test_pdf_extraction_preserves_page_provenance() -> None:
    extracted = CVFileExtractionService().extract(
        filename="cv.pdf",
        content_type="application/pdf",
        content=_pdf_with_pages("First page", "Second page"),
    )
    assert [segment.page_number for segment in extracted.segments] == [1, 2]
    assert [segment.text for segment in extracted.segments] == ["First page", "Second page"]
    assert extracted.provenance.segment_ids == [segment.segment_id for segment in extracted.segments]


def test_invalid_semantic_output_is_rejected_without_persistence(db_session, runtime_snapshot_a) -> None:
    interpreter = SemanticCVInterpreter(InvalidResponseClient(), "test-model")
    service = CVIngestionService(db_session, interpreter=interpreter, runtime_snapshot=runtime_snapshot_a)
    draft = service.upload("user-1", [("cv.md", "text/markdown", b"# Experience\nPython engineer")])
    with pytest.raises(RuntimeError, match="invalid structured data"):
        service.interpret("user-1", draft.id)
    assert db_session.scalars(select(CandidateEvidenceRecord)).all() == []
    assert db_session.scalars(select(CandidateStructuredProfile)).all() == []


def test_semantic_cv_interpreter_uses_provider_strict_schema_and_keeps_pydantic_authority() -> None:
    client = CapturingResponseClient('{"employment":[],"education":[],"credentials":[],"skills":[],"projects":[],"achievements":[],"evidence":[]}')
    document = CVFileExtractionService().extract(
        filename="cv.md", content_type="text/markdown", content=b"# CV\nSource text"
    )
    result = SemanticCVInterpreter(client, "test-model").interpret([document])

    assert result.credentials == []
    assert client.kwargs is not None
    response_format = client.kwargs["text"]["format"]
    assert response_format["type"] == "json_schema"
    assert response_format["strict"] is True
    schema = response_format["schema"]
    assert set(schema["required"]) == set(schema["properties"])
    payload = json.loads(client.kwargs["input"][1]["content"].split("INPUT:\n", 1)[1])
    supplied_document = payload["documents"][0]
    provenance_schema = schema["$defs"]["EvidenceProvenance"]["properties"]
    assert provenance_schema["document_sha256"]["enum"] == [document.provenance.document_sha256]
    assert provenance_schema["segment_ids"]["items"]["enum"] == document.provenance.segment_ids
    assert [segment["segment_id"] for segment in supplied_document["segments"]] == document.provenance.segment_ids
    assert supplied_document["provenance"]["segment_ids"] == provenance_schema["segment_ids"]["items"]["enum"]
    assert "Copy each `segment_ids` value exactly" in client.kwargs["input"][0]["content"]
    encoded_schema = json.dumps(schema)
    assert "certification" in encoded_schema
    assert "professional_registration" in encoded_schema
    assert "source_ref" not in encoded_schema
    assert '"cv"' in encoded_schema
    assert "JSON schema:" not in client.kwargs["input"][1]["content"]

    with pytest.raises(Exception):
        CandidateCVData.model_validate(
            {"credentials": [{"name": "Cert", "credential_type": "unsupported"}]}
        )
    with pytest.raises(Exception):
        CandidateCVData.model_validate(
            {"credentials": [{"name": "Cert", "credential_type": "certification", "extra": "no"}]}
        )
    with pytest.raises(Exception):
        CandidateCVData.model_validate(
            {
                "evidence": [
                    {
                        "evidence_type": "project",
                        "title": "Claim",
                        "text": "Claim text.",
                        "provenance": [{"document_sha256": "a", "source_kind": "confirmed_profile"}],
                    }
                ]
            }
        )


def test_benchmark_markdown_with_unknown_provider_segment_is_rejected_without_promotion(client, db_session, runtime_snapshot_a) -> None:
    content = _benchmark_cv_bytes()
    extracted = CVFileExtractionService().extract(
        filename="issue_254_benchmark_cv.md", content_type="text/markdown", content=content
    )
    invalid_output = _evidence_output(
        extracted.provenance.document_sha256, ["segment-1"]
    )
    provider = CapturingResponseClient(invalid_output.model_dump_json())
    service = CVIngestionService(
        db_session,
        interpreter=SemanticCVInterpreter(provider, "test-model"),
        runtime_snapshot=runtime_snapshot_a,
    )
    app.dependency_overrides[get_user_cv_ingestion_service] = lambda: service
    try:
        headers = _auth(client, "issue-254-unknown@example.com")
        upload = client.post(
            "/api/v1/cv-ingestion/upload",
            headers=headers,
            files=[("files", ("issue_254_benchmark_cv.md", content, "text/markdown"))],
        )
        assert upload.status_code == 201
        draft_id = upload.json()["id"]
        response = client.post(f"/api/v1/cv-ingestion/{draft_id}/interpret", headers=headers)
        assert response.status_code == 422
        assert response.json()["detail"] == "CV evidence provenance must reference supplied source segments."

        request_payload = json.loads(provider.kwargs["input"][1]["content"].split("INPUT:\n", 1)[1])
        request_document = request_payload["documents"][0]
        request_segment_ids = [segment["segment_id"] for segment in request_document["segments"]]
        provenance_schema = provider.kwargs["text"]["format"]["schema"]["$defs"]["EvidenceProvenance"]["properties"]
        assert request_segment_ids == extracted.provenance.segment_ids
        assert request_document["provenance"]["segment_ids"] == request_segment_ids
        assert provenance_schema["document_sha256"]["enum"] == [extracted.provenance.document_sha256]
        assert provenance_schema["segment_ids"]["items"]["enum"] == extracted.provenance.segment_ids
        assert "segment-1" not in provenance_schema["segment_ids"]["items"]["enum"]

        persisted = db_session.get(CandidateCVIngestionDraft, draft_id)
        assert persisted is not None and persisted.state == "uploaded" and persisted.merged_json is None
        assert db_session.scalar(select(CandidateCVReviewBaseline).where(CandidateCVReviewBaseline.draft_id == draft_id)) is None
        assert db_session.scalars(select(CandidateEvidenceRecord)).all() == []
        assert db_session.scalars(select(CandidateStructuredProfile)).all() == []
    finally:
        app.dependency_overrides.pop(get_user_cv_ingestion_service, None)


def test_valid_benchmark_markdown_provenance_is_accepted_and_confirmed(client, db_session, runtime_snapshot_a) -> None:
    content = _benchmark_cv_bytes()
    extracted = CVFileExtractionService().extract(
        filename="issue_254_benchmark_cv.md", content_type="text/markdown", content=content
    )
    interpreter = FixedCVInterpreter(
        _evidence_output(extracted.provenance.document_sha256, [extracted.segments[0].segment_id])
    )
    service = CVIngestionService(db_session, interpreter=interpreter, runtime_snapshot=runtime_snapshot_a)
    app.dependency_overrides[get_user_cv_ingestion_service] = lambda: service
    try:
        headers = _auth(client, "issue-254-valid@example.com")
        upload = client.post(
            "/api/v1/cv-ingestion/upload",
            headers=headers,
            files=[("files", ("issue_254_benchmark_cv.md", content, "text/markdown"))],
        )
        draft_id = upload.json()["id"]
        interpreted = client.post(f"/api/v1/cv-ingestion/{draft_id}/interpret", headers=headers)
        assert interpreted.status_code == 200
        assert interpreted.json()["state"] == "review_ready"
        provenance = interpreted.json()["merged"]["evidence"][0]["provenance"][0]
        assert provenance == {
            "document_sha256": extracted.provenance.document_sha256,
            "segment_ids": [extracted.segments[0].segment_id],
            "source_kind": "cv",
        }
        confirmed = client.post(f"/api/v1/cv-ingestion/{draft_id}/confirm", headers=headers)
        assert confirmed.status_code == 200
        record = db_session.scalar(select(CandidateEvidenceRecord))
        assert record is not None
        assert json.loads(record.provenance_json) == [{**provenance, "source_ref": None}]
    finally:
        app.dependency_overrides.pop(get_user_cv_ingestion_service, None)


def test_mixed_valid_and_unknown_provider_provenance_rejects_whole_interpretation(client, db_session, runtime_snapshot_a) -> None:
    content = _benchmark_cv_bytes()
    extracted = CVFileExtractionService().extract(
        filename="issue_254_benchmark_cv.md", content_type="text/markdown", content=content
    )
    output = _evidence_output(extracted.provenance.document_sha256, [extracted.segments[0].segment_id])
    output.evidence.append(_evidence_output(
        extracted.provenance.document_sha256, ["unknown-segment"], title="Unknown source claim"
    ).evidence[0])
    service = CVIngestionService(
        db_session, interpreter=FixedCVInterpreter(output), runtime_snapshot=runtime_snapshot_a
    )
    app.dependency_overrides[get_user_cv_ingestion_service] = lambda: service
    try:
        headers = _auth(client, "issue-254-mixed@example.com")
        uploaded = client.post(
            "/api/v1/cv-ingestion/upload",
            headers=headers,
            files=[("files", ("issue_254_benchmark_cv.md", content, "text/markdown"))],
        )
        draft_id = uploaded.json()["id"]
        response = client.post(f"/api/v1/cv-ingestion/{draft_id}/interpret", headers=headers)
        assert response.status_code == 422
        assert db_session.get(CandidateCVIngestionDraft, draft_id).state == "uploaded"
        assert db_session.scalars(select(CandidateEvidenceRecord)).all() == []
        assert db_session.scalars(select(CandidateStructuredProfile)).all() == []
    finally:
        app.dependency_overrides.pop(get_user_cv_ingestion_service, None)


def test_provenance_cannot_pair_a_valid_segment_with_the_wrong_document() -> None:
    extractor = CVFileExtractionService()
    first = extractor.extract(filename="first.md", content_type="text/markdown", content=b"# First\nSource A")
    second = extractor.extract(filename="second.md", content_type="text/markdown", content=b"# Second\nSource B")
    cross_document = _evidence_output(first.provenance.document_sha256, [second.segments[0].segment_id])
    with pytest.raises(ValueError, match="supplied source segments"):
        CVIngestionService._enrich_provenance(
            cross_document, [first, second], allow_missing_provenance=False
        )


def test_interpret_state_machine_rejects_repeat_and_preserves_confirmed_data(db_session, runtime_snapshot_a) -> None:
    service = CVIngestionService(db_session, interpreter=FakeInterpreter(), runtime_snapshot=runtime_snapshot_a)
    draft = service.upload("user-1", [("cv.md", "text/markdown", b"# Experience\nPython engineer")])
    reviewed = service.interpret("user-1", draft.id)
    with pytest.raises(ValueError, match="Only an uploaded"):
        service.interpret("user-1", draft.id)
    assert service.read("user-1", draft.id).state == "review_ready"
    assert service.confirm("user-1", draft.id) == 0
    confirmed = service.read("user-1", draft.id)
    persisted = db_session.scalar(select(CandidateStructuredProfile).where(CandidateStructuredProfile.user_id == "user-1"))
    assert confirmed.state == "confirmed"
    assert persisted is not None
    assert json.loads(persisted.structured_json) == reviewed.merged.model_dump(mode="json")
    with pytest.raises(ValueError, match="Only an uploaded"):
        service.interpret("user-1", draft.id)
    assert service.read("user-1", draft.id).state == "confirmed"
    assert db_session.scalar(select(CandidateStructuredProfile).where(CandidateStructuredProfile.user_id == "user-1")).structured_json == persisted.structured_json


def test_review_patch_is_owned_and_persists_only_after_confirm(client, db_session) -> None:
    service = CVIngestionService(db_session, interpreter=FakeInterpreter())
    app.dependency_overrides[get_user_cv_ingestion_service] = lambda: service
    try:
        owner = _auth(client, "draft-owner@example.com")
        other = _auth(client, "draft-other@example.com")
        upload = client.post("/api/v1/cv-ingestion/upload", headers=owner, files=[("files", ("cv.json", json.dumps(_json_cv()), "application/json"))])
        draft_id = upload.json()["id"]
        assert client.patch(f"/api/v1/cv-ingestion/{draft_id}", headers=owner, json=_json_cv()).status_code == 409
        interpreted = client.post(f"/api/v1/cv-ingestion/{draft_id}/interpret", headers=owner)
        assert interpreted.status_code == 200
        # Review clients round-trip the authoritative interpreted form, including
        # source provenance; only ordinary structured fields change here.
        corrected = interpreted.json()["merged"]
        corrected["skills"] = [{"name": "Rust", "category": "technical"}]
        owner_edit = client.patch(f"/api/v1/cv-ingestion/{draft_id}", headers=owner, json=corrected)
        assert owner_edit.status_code == 200
        assert owner_edit.json()["merged"]["skills"] == [{"name": "Rust", "category": "technical"}]
        assert client.patch(f"/api/v1/cv-ingestion/{draft_id}", headers=other, json=corrected).status_code == 404
        assert db_session.scalars(select(CandidateEvidenceRecord)).all() == []
        assert db_session.scalars(select(CandidateStructuredProfile)).all() == []
        assert client.post(f"/api/v1/cv-ingestion/{draft_id}/confirm", headers=owner).status_code == 200
        owner_id = db_session.scalar(select(User.id).where(User.email == "draft-owner@example.com"))
        reader = CanonicalCandidateReadService(db_session)
        context = reader.candidate_context(reader.read(owner_id), require_complete_evidence=True)
        assert context is not None and context.skills_text == "Rust"
    finally:
        app.dependency_overrides.pop(get_user_cv_ingestion_service, None)


def test_upload_batch_and_media_validation(db_session) -> None:
    service = CVIngestionService(db_session, interpreter=FakeInterpreter())
    with pytest.raises(ValueError, match="at most"):
        service.upload("user-1", [(f"cv-{index}.md", "text/markdown", b"CV") for index in range(6)])
    with pytest.raises(ValueError, match="media type"):
        service.upload("user-1", [("cv.md", "application/pdf", b"CV")])
    with pytest.raises(ValueError, match="5 MB"):
        service.upload("user-1", [("cv.md", "text/plain", b"x" * (5 * 1024 * 1024 + 1))])


def test_review_evidence_baseline_is_immutable_subset_and_structured_fields_stay_editable(db_session) -> None:
    service = CVIngestionService(db_session, interpreter=FakeInterpreter())
    draft = service.upload("user-1", [("cv.json", "application/json", json.dumps(_json_cv()).encode())])
    reviewed = service.interpret("user-1", draft.id)
    baseline = db_session.scalar(
        select(CandidateCVReviewBaseline).where(CandidateCVReviewBaseline.draft_id == draft.id)
    )
    assert baseline is not None
    original = reviewed.merged.model_dump(mode="json")

    reordered = CandidateCVData.model_validate({**original, "evidence": list(reversed(original["evidence"]))})
    service.edit_review("user-1", draft.id, reordered)
    excluded = CandidateCVData.model_validate({**original, "evidence": []})
    service.edit_review("user-1", draft.id, excluded)
    structured = CandidateCVData.model_validate({**original, "employment": [{**original["employment"][0], "title": "Principal Engineer"}]})
    service.edit_review("user-1", draft.id, structured)
    assert service.read("user-1", draft.id).merged.employment[0].title == "Principal Engineer"

    for mutated in (
        {**original["evidence"][0], "title": "Invented title"},
        {**original["evidence"][0], "text": "Invented text"},
        {**original["evidence"][0], "skills": ["Invented"]},
        {**original["evidence"][0], "provenance": [{"document_sha256": "invented", "segment_ids": ["invented:1"]}]},
    ):
        with pytest.raises(ValueError, match="Career Evidence"):
            service.edit_review("user-1", draft.id, CandidateCVData.model_validate({**original, "evidence": [mutated]}))
    with pytest.raises(ValueError, match="Career Evidence"):
        service.edit_review("user-1", draft.id, CandidateCVData.model_validate({**original, "evidence": [original["evidence"][0], original["evidence"][0]]}))
    assert db_session.scalar(select(CandidateCVReviewBaseline).where(CandidateCVReviewBaseline.draft_id == draft.id)).evidence_json == baseline.evidence_json


def test_legacy_review_draft_snapshots_once_before_first_edit(db_session) -> None:
    service = CVIngestionService(db_session, interpreter=FakeInterpreter())
    draft = service.upload("user-1", [("cv.json", "application/json", json.dumps(_json_cv()).encode())])
    reviewed = service.interpret("user-1", draft.id)
    baseline = db_session.scalar(select(CandidateCVReviewBaseline).where(CandidateCVReviewBaseline.draft_id == draft.id))
    db_session.delete(baseline)
    db_session.commit()
    original = reviewed.merged.model_dump(mode="json")
    service.edit_review("user-1", draft.id, CandidateCVData.model_validate({**original, "evidence": []}))
    created = db_session.scalar(select(CandidateCVReviewBaseline).where(CandidateCVReviewBaseline.draft_id == draft.id))
    assert created is not None
    with pytest.raises(ValueError, match="Career Evidence"):
        service.edit_review("user-1", draft.id, CandidateCVData.model_validate({**original, "evidence": [{**original["evidence"][0], "text": "changed"}]}))



def test_duplicate_baseline_race_recovers_the_single_immutable_row(db_session, monkeypatch) -> None:
    """Exercise the IntegrityError recovery path that closes a duplicate-create race."""
    service = CVIngestionService(db_session, interpreter=FakeInterpreter())
    uploaded = service.upload("user-race", [("cv.json", "application/json", json.dumps(_json_cv()).encode())])
    reviewed = service.interpret("user-race", uploaded.id)
    existing = db_session.scalar(
        select(CandidateCVReviewBaseline).where(CandidateCVReviewBaseline.draft_id == uploaded.id)
    )
    draft = db_session.scalar(
        select(CandidateCVIngestionDraft).where(CandidateCVIngestionDraft.id == uploaded.id)
    )
    assert existing is not None and draft is not None and reviewed.merged is not None

    real_scalar = db_session.scalar
    baseline_lookup_calls = 0

    def racing_scalar(statement, *args, **kwargs):
        nonlocal baseline_lookup_calls
        rendered = str(statement)
        if "candidate_cv_review_baselines" in rendered:
            baseline_lookup_calls += 1
            if baseline_lookup_calls == 1:
                # Simulate this session checking just before another writer commits
                # the row that is already present in the database.
                return None
        return real_scalar(statement, *args, **kwargs)

    monkeypatch.setattr(db_session, "scalar", racing_scalar)
    recovered = service._create_baseline_once(draft, reviewed.merged.evidence)

    assert recovered.id == existing.id
    rows = db_session.scalars(
        select(CandidateCVReviewBaseline).where(CandidateCVReviewBaseline.draft_id == uploaded.id)
    ).all()
    assert [row.id for row in rows] == [existing.id]
