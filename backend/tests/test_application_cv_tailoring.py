"""Provider-free regressions for profile-first application CV tailoring."""

import json
import re
import sqlite3
from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.agents.application_cv_tailoring import OpenAIApplicationCVTailoringAgent
from app.core.config import Settings
from app.models.application_cv_tailoring import ApplicationCVTailoringSession
from app.models.application_preparation import ApplicationPreparation
from app.models.candidate_profile import CandidateProfile
from app.models.candidate_cv_ingestion import CandidateCVIngestionDraft, CandidateStructuredProfile
from app.models.cv_visibility_assessment import CVVisibilityAssessment
from app.models.discovered_job import DiscoveredJob
from app.models.user import User
from app.api.deps import get_application_cv_tailoring_service
from app.core.security import create_access_token
from app.main import app as fastapi_app
from app.models.user_job_discovery import UserJobEvaluation
from app.providers.llm import OpenAISemanticLLM, SemanticOutputError, SemanticResponseClient
from app.schemas.ai_settings import SemanticOperation
from app.schemas.application_cv_tailoring import (
    TailoringAction, TailoringApproveWrite, TailoringDraftGenerateRequest,
    TailoringDraftReviewWrite, TailoringRecommendation, TailoringStartRequest,
    TailoringStrategy, TailoringStrategyReview, TailoringStrategyReviewWrite,
)
from app.schemas.application_preparation import (
    ApplicationIdentitySnapshot, ApplicationSourceRef, ApplicationTargetSnapshot,
    ApplicationTargetKind, CVWritingDraft, TailoredBullet, TailoredRoleDraft,
)
from app.schemas.assessment import FitAssessment
from app.schemas.candidate import CandidateContext
from app.schemas.cv_ingestion import CandidateCVData, CareerEvidenceDraft, Employment, EvidenceProvenance, Project, Skill
from app.schemas.career_assessment import AlignmentConfidence, CareerAssessment
from app.schemas.job import JobProfile, JobRequirement
from app.schemas.job_ranking import (
    JobArchetype, JobArchetypeAssessment, JobRelevanceAssessment,
    PostingLegitimacy, PostingLegitimacyAssessment, RankedJobOpportunity,
)
from app.schemas.matching import EvidenceRef, EvidenceSourceType, RequirementMatch
from app.schemas.recommendation import Recommendation, RecommendationAssessment
from app.services.application_cv_tailoring_service import (
    ApplicationCVTailoringError, ApplicationCVTailoringService, _Authority, _runtime_fingerprint,
)
from app.services.application_preparation_service import (
    ApplicationPreparationReadService, ApplicationPreparationService,
)
from app.services.user_job_discovery_service import UserJobDiscoveryService
from app.services.active_candidate_evidence import ActiveCandidateEvidenceResolver
from app.services.canonical_candidate_read_service import CanonicalCandidateReadService
from app.services.user_job_discovery_service import evaluation_contract_fingerprint_for_runtime


def _make_authority(db_session, runtime_snapshot_a, user_id: str, job_id: str) -> _Authority:
    job = DiscoveredJob(
        id=job_id, identity_key=f"provider:{job_id}", source="synthetic", title="AI systems engineer",
        location="London", url="https://jobs.example.test/role", description="Build reliable systems.",
        content_hash="a" * 64, state="unchanged", verification_status="verified",
        detail_authority="verified_employer_detail", first_seen_at=datetime.now(timezone.utc),
        last_seen_at=datetime.now(timezone.utc), last_changed_at=datetime.now(timezone.utc),
    )
    db_session.add(job)
    evaluation = UserJobEvaluation(
        id="evaluation-current", user_id=user_id, discovered_job_id=job_id,
        job_content_hash=job.content_hash, candidate_evaluation_fingerprint="candidate-fingerprint",
        evaluation_contract_fingerprint="evaluation-contract", job_snapshot_json="{}", evaluation_json="{}",
    )
    db_session.add(evaluation)
    db_session.commit()
    skill_ref = ApplicationSourceRef(source_type="skills", source_ref="skill:0")
    requirement = JobRequirement(text="Build Python services", importance="essential", category="technical")
    match = RequirementMatch(
        requirement_index=0, requirement=requirement, match_type="demonstrated", score=0.9,
        evidence_ids=["confirmed-python"], evidence_refs=[EvidenceRef(source_type="skills", source_ref="skill:0", value="Python")], reasoning="Confirmed profile evidence.",
    )
    target = ApplicationTargetSnapshot(
        source_kind=ApplicationTargetKind.DISCOVERED_JOB, canonical_discovered_job_id=job_id,
        public_url=job.url, title=job.title, company="Synthetic Co", location="London",
        job_profile=JobProfile(title=job.title, requirements=[requirement]), requirement_matches=[match],
        job_content_hash=job.content_hash,
    )
    structured = CandidateCVData(
        skills=[Skill(name="Python")],
        employment=[Employment(employer="Example Co", title="Engineer", start_date="2020", end_date="2024", description="Built data services.")],
    )
    sources = {
        ("skills", "skill:0"): "Python",
        ("employment", "employment:0"): "Engineer at Example Co, 2020 to 2024. Built data services.",
    }
    return _Authority(
        job=job, evaluation=evaluation, opportunity=None, target=target,
        identity=ApplicationIdentitySnapshot(display_name="Synthetic Person", email="person@example.test"),
        context=CandidateContext(profile_text="Confirmed candidate profile", skills_text="Python"),
        structured=structured, sources=sources, candidate_fingerprint="full-profile-fingerprint",
        runtime_fingerprint=_runtime_fingerprint(runtime_snapshot_a),
    )


class _Agent:
    def __init__(self, strategy=None, draft=None, error=None):
        self.strategy = strategy
        self.draft = draft
        self.error = error
        self.strategy_calls = 0
        self.draft_calls = 0
        self.strategy_contexts = []
        self.draft_contexts = []

    def recommend_strategy(self, _context):
        self.strategy_calls += 1
        self.strategy_contexts.append(_context)
        if self.error:
            raise self.error
        return self.strategy

    def draft_tailored_cv(self, _context):
        self.draft_calls += 1
        self.draft_contexts.append(_context)
        if self.error:
            raise self.error
        return self.draft


class _RecordingClient:
    def __init__(self, response):
        self.response = response
        self.calls = []
        self.responses = self

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(output_text=self.response)


def _mock_openai_semantic_adapter(monkeypatch, output_text, operation):
    """Capture the provider request after the semantic compatibility adapter."""
    calls = []

    class _Responses:
        def create(self, **kwargs):
            calls.append(kwargs)
            return SimpleNamespace(output_text=output_text)

    client = SimpleNamespace(responses=_Responses())
    monkeypatch.setattr("app.providers.llm.create_traced_openai_client", lambda **_kwargs: client)
    semantic_client = SemanticResponseClient(
        OpenAISemanticLLM(api_key="synthetic-test-key"), operation=operation,
    )
    return semantic_client, calls


def _run_draft_through_semantic_adapter(monkeypatch, context, draft):
    semantic_client, calls = _mock_openai_semantic_adapter(
        monkeypatch, json.dumps(draft.model_dump(mode="json")), "application_cv_tailoring_draft",
    )
    agent = OpenAIApplicationCVTailoringAgent(
        strategy_client=semantic_client, draft_client=semantic_client, model="gpt-5.6-luna",
    )
    result = agent.draft_tailored_cv(context)
    assert len(calls) == 1
    return result, calls[0]


def _run_strategy_through_semantic_adapter(monkeypatch, context, response_json):
    semantic_client, calls = _mock_openai_semantic_adapter(
        monkeypatch, json.dumps(response_json), "application_cv_tailoring_strategy",
    )
    agent = OpenAIApplicationCVTailoringAgent(
        strategy_client=semantic_client, draft_client=semantic_client, model="gpt-5.6-luna",
    )
    result = agent.recommend_strategy(context)
    assert len(calls) == 1
    return result, calls[0], calls


def _assert_provider_schema_has_no_array_bounds(schema):
    if isinstance(schema, dict):
        assert "maxItems" not in schema
        assert "minItems" not in schema
        for value in schema.values():
            _assert_provider_schema_has_no_array_bounds(value)
    elif isinstance(schema, list):
        for value in schema:
            _assert_provider_schema_has_no_array_bounds(value)


def _strategy() -> TailoringStrategy:
    return TailoringStrategy(
        positioning="Evidence-led Python delivery.", summary_emphasis="Python services and delivery.",
        recommendations=[TailoringRecommendation(
            recommendation_id="requirement-0", requirement_index=0, action=TailoringAction.FEATURE,
            rationale="Confirmed Python skills and delivery experience support this requirement.",
            source_refs=[ApplicationSourceRef(source_type="skills", source_ref="skill:0"), ApplicationSourceRef(source_type="employment", source_ref="employment:0")],
        )],
        gaps_to_clarify=[], not_to_claim=["Do not claim unsupported cloud certifications."],
    )


def _draft() -> CVWritingDraft:
    skill = ApplicationSourceRef(source_type="skills", source_ref="skill:0")
    role = ApplicationSourceRef(source_type="employment", source_ref="employment:0")
    return CVWritingDraft(
        professional_summary="Engineer building Python services.", summary_source_refs=[skill], key_skills=["Python"],
        role_drafts=[TailoredRoleDraft(employment_index=0, bullets=[TailoredBullet(
            text="Built data services.", source_refs=[skill, role], priority=70,
        )])],
    )


def _service(db_session, runtime_snapshot_a, authority, agent):
    user_job_service = UserJobDiscoveryService(db_session, runtime_snapshot=runtime_snapshot_a)
    application_service = ApplicationPreparationService(
        db_session, user_discovery=user_job_service, graph_factory=lambda: None,
        drafting_agent_factory=lambda: None, settings=Settings(_env_file=None), runtime_snapshot=runtime_snapshot_a,
    )
    original_context = application_service._draft_context
    def draft_context(target, catalog, structured):
        return {
            "job": {"title": target.title, "requirements": [item.requirement.text for item in target.requirement_matches]},
            "sources": [{"source_type": kind, "source_ref": ref, "text": value} for (kind, ref), value in catalog.items()],
            "allowed_skills": [item.name for item in structured.skills],
            "employment": [{"employment_index": 0, "required_anchor_ref": {"source_type": "employment", "source_ref": "employment:0"}}],
            "projects": [],
        }
    application_service._draft_context = draft_context
    application_service._bounded_sources = lambda _context, _structured, _matches: authority.sources
    service = ApplicationCVTailoringService(
        db_session, runtime_snapshot=runtime_snapshot_a, application_service=application_service,
        agent_factory=lambda: agent, settings=Settings(_env_file=None),
    )
    service._authority = lambda _user_id, _job_id, runtime=None: authority
    return service


def test_generated_strategy_schema_is_strict_and_scoped_to_current_profile_and_requirements():
    schema = OpenAIApplicationCVTailoringAgent._strategy_schema({
        "sources": [{"source_type": "skills", "source_ref": "skill:0"}],
        "requirements": [{"requirement_index": 2}, {"requirement_index": 7}],
    })
    assert schema["additionalProperties"] is False
    assert schema["$defs"]["TailoringRecommendation"]["properties"]["requirement_index"]["enum"] == [2, 7]
    assert schema["$defs"]["ApplicationSourceRef"]["anyOf"][0]["properties"]["source_ref"]["enum"] == ["skill:0"]

    client = _RecordingClient(json.dumps(_strategy().model_dump(mode="json")))
    agent = OpenAIApplicationCVTailoringAgent(strategy_client=client, draft_client=client, model="mock-model")
    result = agent.recommend_strategy({
        "sources": [{"source_type": "skills", "source_ref": "skill:0"}, {"source_type": "employment", "source_ref": "employment:0"}],
        "requirements": [{"requirement_index": 0}],
    })
    assert result.recommendations[0].recommendation_id == "requirement-0"
    assert len(client.calls) == 1
    assert client.calls[0]["text"]["format"]["strict"] is True

    draft_client = _RecordingClient(json.dumps(_draft().model_dump(mode="json")))
    draft_agent = OpenAIApplicationCVTailoringAgent(strategy_client=client, draft_client=draft_client, model="mock-model")
    draft_agent.draft_tailored_cv({
        "sources": [{"source_type": "skills", "source_ref": "skill:0"}, {"source_type": "employment", "source_ref": "employment:0"}],
        "allowed_skills": ["Python"], "employment": [{"employment_index": 0}], "projects": [],
    })
    draft_schema = draft_client.calls[0]["text"]["format"]["schema"]
    assert draft_schema["$defs"]["TailoredRoleDraft"]["properties"]["employment_index"]["enum"] == [0]
    assert draft_schema["properties"]["key_skills"]["items"]["enum"] == ["Python"]
    assert len(draft_client.calls) == 1


def test_v2b_strategy_schema_is_normalized_after_scope_and_enum_restrictions(monkeypatch):
    context = {
        "sources": [
            {"source_type": "career_evidence", "source_ref": "evidence:research", "text": "Built a research prototype."},
            {"source_type": "skills", "source_ref": "skills:0", "text": "Python"},
        ],
        "requirements": [{"requirement_index": 2}, {"requirement_index": 7}],
    }
    result_payload = TailoringStrategy(
        positioning="Evidence-led applied AI engineering.", summary_emphasis="Research and Python delivery.",
        recommendations=[
            TailoringRecommendation(
                recommendation_id="requirement-2", requirement_index=2, action=TailoringAction.FEATURE,
                rationale="Confirmed research evidence supports this requirement.",
                source_refs=[ApplicationSourceRef(source_type="career_evidence", source_ref="evidence:research")],
            ),
            TailoringRecommendation(
                recommendation_id="requirement-7", requirement_index=7, action=TailoringAction.FEATURE,
                rationale="Confirmed Python evidence supports this requirement.",
                source_refs=[ApplicationSourceRef(source_type="skills", source_ref="skills:0")],
            ),
        ], gaps_to_clarify=[], not_to_claim=["Do not claim production ownership."],
    ).model_dump(mode="json")

    result, request, calls = _run_strategy_through_semantic_adapter(monkeypatch, context, result_payload)
    format_request = request["text"]["format"]
    schema = format_request["schema"]

    assert format_request["strict"] is True
    assert format_request["name"] == "application_cv_tailoring_strategy"
    _assert_provider_schema_has_no_array_bounds(schema)
    recommendation_properties = schema["$defs"]["TailoringRecommendation"]["properties"]
    assert recommendation_properties["requirement_index"]["enum"] == [2, 7]
    assert recommendation_properties["recommendation_id"]["enum"] == ["requirement-2", "requirement-7"]
    source_pairs = {
        (item["properties"]["source_type"]["enum"][0], item["properties"]["source_ref"]["enum"][0])
        for item in schema["$defs"]["ApplicationSourceRef"]["anyOf"]
    }
    assert source_pairs == {("career_evidence", "evidence:research"), ("skills", "skills:0")}
    assert [(item.requirement_index, item.recommendation_id) for item in result.recommendations] == [
        (2, "requirement-2"), (7, "requirement-7"),
    ]
    assert len(calls) == 1


def test_v2b_strategy_pydantic_bounds_reject_oversized_model_recommendations(monkeypatch):
    context = {
        "sources": [{"source_type": "career_evidence", "source_ref": "evidence:confirmed", "text": "Confirmed experience."}],
        "requirements": [{"requirement_index": index} for index in range(6)],
    }
    response_payload = {
        "positioning": "Evidence-led candidate.", "summary_emphasis": "Confirmed experience.",
        "recommendations": [
            {
                "recommendation_id": f"requirement-{index}", "requirement_index": index,
                "action": "feature", "rationale": "This confirmed evidence supports the requirement.",
                "suggested_emphasis": None,
                "source_refs": [{"source_type": "career_evidence", "source_ref": "evidence:confirmed"}],
            }
            for index in range(6)
        ],
        "gaps_to_clarify": [], "not_to_claim": [],
    }
    semantic_client, calls = _mock_openai_semantic_adapter(
        monkeypatch, json.dumps(response_payload), "application_cv_tailoring_strategy",
    )
    agent = OpenAIApplicationCVTailoringAgent(
        strategy_client=semantic_client, draft_client=semantic_client, model="gpt-5.6-luna",
    )

    with pytest.raises(SemanticOutputError, match="invalid structured output"):
        agent.recommend_strategy(context)

    assert len(calls) == 1
    schema = calls[0]["text"]["format"]["schema"]
    _assert_provider_schema_has_no_array_bounds(schema)
    assert schema["$defs"]["TailoringRecommendation"]["properties"]["requirement_index"]["enum"] == list(range(6))


def test_v2b_semantic_provider_schema_accepts_career_evidence_only_with_empty_sections(monkeypatch):
    evidence = ApplicationSourceRef(source_type="career_evidence", source_ref="evidence:research")
    context = {
        "sources": [{"source_type": "career_evidence", "source_ref": "evidence:research", "text": "Built a retrieval research prototype."}],
        "selected_support_refs": [{"source_type": "career_evidence", "source_ref": "evidence:research"}],
        "allowed_skills": [], "employment": [], "projects": [],
    }
    expected = CVWritingDraft(
        professional_summary="Engineer with confirmed retrieval research experience.",
        summary_source_refs=[evidence], key_skills=[], role_drafts=[], project_drafts=[],
    )

    result, request = _run_draft_through_semantic_adapter(monkeypatch, context, expected)
    format_request = request["text"]["format"]
    schema = format_request["schema"]

    assert format_request["strict"] is True
    _assert_provider_schema_has_no_array_bounds(schema)
    assert result.key_skills == result.role_drafts == result.project_drafts == []
    ApplicationCVTailoringService._validate_draft_strategy_refs(
        result, {("career_evidence", "evidence:research"): "Built a retrieval research prototype."}, {}, CandidateCVData(),
    )


def test_v2b_semantic_provider_schema_allows_empty_skills_and_projects_with_selected_role(monkeypatch):
    evidence = ApplicationSourceRef(source_type="career_evidence", source_ref="evidence:employment")
    role_anchor = ApplicationSourceRef(source_type="employment", source_ref="employment:0")
    context = {
        "sources": [
            {"source_type": "career_evidence", "source_ref": "evidence:employment", "text": "Built data services."},
            {"source_type": "employment", "source_ref": "employment:0", "text": "Engineer at Example Co, 2020 to 2024"},
        ],
        "selected_support_refs": [{"source_type": "career_evidence", "source_ref": "evidence:employment"}],
        "structural_attribution_refs": [{"source_type": "employment", "source_ref": "employment:0"}],
        "allowed_skills": [], "employment": [{"employment_index": 0}], "projects": [],
    }
    expected = CVWritingDraft(
        professional_summary="Engineer with confirmed data services experience.", summary_source_refs=[evidence], key_skills=[],
        role_drafts=[TailoredRoleDraft(employment_index=0, bullets=[TailoredBullet(
            text="Built data services.", source_refs=[evidence, role_anchor], priority=70,
        )])],
        project_drafts=[],
    )

    result, request = _run_draft_through_semantic_adapter(monkeypatch, context, expected)
    schema = request["text"]["format"]["schema"]
    _assert_provider_schema_has_no_array_bounds(schema)
    assert schema["$defs"]["TailoredRoleDraft"]["properties"]["employment_index"]["enum"] == [0]
    assert result.key_skills == [] and result.project_drafts == []
    ApplicationCVTailoringService._validate_draft_strategy_refs(
        result,
        {("career_evidence", "evidence:employment"): "Built data services."},
        {("employment", "employment:0"): "Engineer at Example Co, 2020 to 2024"},
        CandidateCVData(employment=[Employment(employer="Example Co", title="Engineer", start_date="2020", end_date="2024")]),
    )


@pytest.mark.parametrize("kind", ["skill", "role", "project"])
def test_v2b_application_validation_rejects_content_outside_empty_selected_sections(kind):
    evidence = ApplicationSourceRef(source_type="career_evidence", source_ref="evidence:selected")
    values = {
        "professional_summary": "Engineer with selected evidence.",
        "summary_source_refs": [evidence],
        "key_skills": [], "role_drafts": [], "project_drafts": [],
    }
    if kind == "skill":
        values["key_skills"] = ["Unselected PyTorch"]
    elif kind == "role":
        values["role_drafts"] = [TailoredRoleDraft(employment_index=0, bullets=[TailoredBullet(
            text="Unselected employment claim.",
            source_refs=[evidence, ApplicationSourceRef(source_type="employment", source_ref="employment:0")], priority=70,
        )])]
    else:
        values["project_drafts"] = [{
            "project_index": 0, "text": "Unselected project claim.",
            "source_refs": [evidence, ApplicationSourceRef(source_type="project", source_ref="project:0")], "priority": 70,
        }]
    draft = CVWritingDraft.model_validate(values)

    with pytest.raises(ValueError):
        ApplicationCVTailoringService._validate_draft_strategy_refs(
            draft, {("career_evidence", "evidence:selected"): "Selected career evidence."}, {}, CandidateCVData(),
        )


def test_sqlite_migration_is_additive_and_idempotent_for_existing_application_data():
    sql_path = Path(__file__).parents[1] / "migrations" / "20261012_application_cv_tailoring_sqlite.sql"
    connection = sqlite3.connect(":memory:")
    connection.executescript("CREATE TABLE users (id VARCHAR(36) PRIMARY KEY); CREATE TABLE application_preparations (id VARCHAR(36) PRIMARY KEY);")
    connection.execute("INSERT INTO users (id) VALUES ('user-existing')")
    connection.execute("INSERT INTO application_preparations (id) VALUES ('preparation-existing')")
    sql = sql_path.read_text(encoding="utf-8")
    connection.executescript(sql)
    connection.executescript(sql)
    assert connection.execute("SELECT id FROM users").fetchall() == [("user-existing",)]
    assert connection.execute("SELECT id FROM application_preparations").fetchall() == [("preparation-existing",)]
    columns = {row[1] for row in connection.execute("PRAGMA table_info(application_cv_tailoring_sessions)")}
    assert {"profile_outline_json", "strategy_history_json", "draft_history_json", "approved_preparation_id"} <= columns
    assert connection.execute("SELECT COUNT(*) FROM application_cv_tailoring_sessions").fetchone() == (0,)
    connection.close()


def test_profile_first_strategy_review_draft_revisions_and_approval_are_durable_and_user_scoped(db_session, runtime_snapshot_a):
    user = User(id="tailor-owner", email="owner@example.test", password_hash="safe")
    other = User(id="tailor-other", email="other@example.test", password_hash="safe")
    db_session.add_all([user, other]); db_session.commit()
    authority = _make_authority(db_session, runtime_snapshot_a, user.id, "tailor-job")
    agent = _Agent(strategy=_strategy(), draft=_draft())
    service = _service(db_session, runtime_snapshot_a, authority, agent)

    # No uploaded source CV or V2A row exists; canonical Profile evidence is sufficient.
    started = service.start(user.id, TailoringStartRequest(
        client_request_id="11111111-1111-4111-8111-111111111111", discovered_job_id=authority.job.id, target_pages=2,
    ))
    assert started.status == "strategy_review", started.failure_code
    assert agent.strategy_calls == 1
    assert service.get_by_request(user.id, started.client_request_id).id == started.id
    with pytest.raises(LookupError):
        service.get(other.id, started.id)

    reviewed = service.review_strategy(user.id, started.id, TailoringStrategyReviewWrite(
        expected_revision=1, accepted=True,
        choices=[TailoringStrategyReview(recommendation_id="requirement-0", decision="keep", order=0)],
    ))
    assert reviewed.status == "strategy_accepted"
    draft = service.generate_draft(user.id, started.id, TailoringDraftGenerateRequest(
        client_request_id="22222222-2222-4222-8222-222222222222", expected_strategy_revision=2,
    ))
    assert draft.status == "draft_review"
    assert agent.draft_calls == 1

    edited_draft = draft.draft.model_copy(update={"professional_summary": "Engineer delivering Python services."})
    saved = service.review_draft(user.id, started.id, TailoringDraftReviewWrite(expected_revision=1, draft=edited_draft))
    assert saved.draft_revision == 2
    assert saved.draft_claim_review_status == "user_edits_unverified"
    assert saved.draft_claim_review_acknowledged is False
    assert len(json.loads(db_session.get(ApplicationCVTailoringSession, started.id).draft_history_json)) == 2
    with pytest.raises(ApplicationCVTailoringError, match="Edited factual claims"):
        service.approve(user.id, started.id, TailoringApproveWrite(expected_strategy_revision=2, expected_draft_revision=2))

    acknowledged = service.review_draft(user.id, started.id, TailoringDraftReviewWrite(
        expected_revision=2, draft=saved.draft, acknowledge_unverified_claims=True,
    ))
    assert acknowledged.draft_claim_review_status == "user_edits_acknowledged_unverified"
    assert acknowledged.draft_claim_review_acknowledged is True

    approved = service.approve(user.id, started.id, TailoringApproveWrite(expected_strategy_revision=2, expected_draft_revision=3))
    assert approved.status == "approved"
    assert approved.approved_preparation_id
    assert db_session.query(ApplicationPreparation).filter_by(user_id=user.id).count() == 1
    final = ApplicationPreparationReadService(db_session).get(user.id, approved.approved_preparation_id)
    assert final.result.cv.professional_summary == "Engineer delivering Python services."
    assert final.result.tailoring_review.draft_claim_status == "user_edits_acknowledged_unverified"
    assert final.result.tailoring_review.draft_user_decision == "acknowledged_unverified_user_edits"
    assert final.target.title == authority.target.title
    assert db_session.query(ApplicationCVTailoringSession).filter_by(user_id=other.id).count() == 0


def test_real_canonical_profile_and_current_evaluation_are_sufficient_without_cv_upload_or_v2a(db_session, runtime_snapshot_a):
    user = User(id="profile-first-owner", email="profile-first@example.test", password_hash="safe")
    db_session.add(user)
    db_session.add(CandidateProfile(
        user_id=user.id, display_name="Synthetic Candidate", preferred_email=user.email,
        headline="Python engineer", summary="Builds reliable services.",
    ))
    structured = CandidateCVData(
        skills=[Skill(name="Python"), Skill(name="HiddenTool")],
        employment=[Employment(employer="Example Co", title="Engineer", start_date="2021", end_date="2025", description="Built Python data services.")],
        projects=[Project(name="Agentic research prototype", description="Evaluated retrieval workflows.", skills=["Python"])],
        evidence=[CareerEvidenceDraft(
            evidence_type="project", title="Agentic research prototype",
            text="Built and evaluated a retrieval research prototype.", skills=["Python"],
            provenance=[EvidenceProvenance(document_sha256="a" * 64, segment_ids=["project-segment"])],
        )],
    )
    db_session.add(CandidateStructuredProfile(user_id=user.id, structured_json=structured.model_dump_json()))
    job = DiscoveredJob(
        id="real-profile-job", identity_key="provider:real-profile-job", source="synthetic",
        title="Applied AI Engineer", company="Synthetic Co", location="London",
        url="https://jobs.example.test/applied-ai", description="Build reliable Python data services.",
        content_hash="b" * 64, state="unchanged", verification_status="verified",
        detail_authority="verified_employer_detail", first_seen_at=datetime.now(timezone.utc),
        last_seen_at=datetime.now(timezone.utc), last_changed_at=datetime.now(timezone.utc),
    )
    db_session.add(job)
    db_session.commit()

    ActiveCandidateEvidenceResolver(db_session).resolve(user.id, structured)
    db_session.commit()
    reader = CanonicalCandidateReadService(db_session)
    snapshot = reader.read(user.id)
    context = reader.candidate_context(snapshot, require_structured_profile=True, require_complete_evidence=True)
    assert context is not None and context.evidence
    employment_evidence = next(item for item in context.evidence if item.evidence_type == "employment")
    project_evidence = next(item for item in context.evidence if item.evidence_type == "project")
    requirements = [
        JobRequirement(text="Build reliable data services", importance="essential", category="technical"),
        JobRequirement(text="Prototype retrieval research", importance="essential", category="experience"),
        JobRequirement(text="Use Python", importance="desirable", category="technical"),
        JobRequirement(text="Use HiddenTool", importance="desirable", category="technical"),
    ]
    matches = [
        RequirementMatch(
            requirement_index=0, requirement=requirements[0], match_type="demonstrated", score=0.9,
            evidence_ids=[employment_evidence.evidence_id],
            evidence_refs=[EvidenceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref=employment_evidence.evidence_id, value=employment_evidence.text)],
            reasoning="Current confirmed employment evidence.",
        ),
        RequirementMatch(
            requirement_index=1, requirement=requirements[1], match_type="transferable", score=0.8,
            evidence_ids=[project_evidence.evidence_id],
            evidence_refs=[EvidenceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref=project_evidence.evidence_id, value=project_evidence.text)],
            reasoning="Current confirmed project evidence.",
        ),
        RequirementMatch(
            requirement_index=2, requirement=requirements[2], match_type="demonstrated", score=0.9,
            evidence_refs=[EvidenceRef(source_type=EvidenceSourceType.SKILLS, source_ref="skills:0", value="Python")],
            reasoning="Current confirmed Profile skill.",
        ),
        RequirementMatch(
            requirement_index=3, requirement=requirements[3], match_type="demonstrated", score=0.8,
            evidence_refs=[EvidenceRef(source_type=EvidenceSourceType.SKILLS, source_ref="skills:1", value="HiddenTool")],
            reasoning="Current confirmed Profile skill.",
        ),
    ]
    opportunity = RankedJobOpportunity(
        job=UserJobDiscoveryService._listing(job),
        relevance=JobRelevanceAssessment(relevant=True, score=0.9, reasoning="Synthetic relevance."),
        archetype=JobArchetypeAssessment(archetype=JobArchetype.OTHER, reasoning="Synthetic archetype."),
        fit_assessment=FitAssessment(fit_score=78),
        career_assessment=CareerAssessment(career_alignment_score=80, confidence=AlignmentConfidence.HIGH, dimensions=[], reasoning="Synthetic alignment."),
        recommendation_assessment=RecommendationAssessment(
            recommendation=Recommendation.CONSIDER, fit_score=78, career_alignment_score=80,
            career_alignment_confidence=AlignmentConfidence.HIGH, rule_id="synthetic",
            reasoning="Synthetic recommendation.",
        ),
        legitimacy=PostingLegitimacyAssessment(legitimacy=PostingLegitimacy.HIGH_CONFIDENCE, reasoning="Synthetic listing."),
        rank=1, job_profile=JobProfile(title=job.title, company=job.company, requirements=requirements),
        requirement_matches=matches,
    )
    db_session.add(UserJobEvaluation(
        id="profile-first-evaluation", user_id=user.id, discovered_job_id=job.id,
        job_content_hash=job.content_hash,
        candidate_evaluation_fingerprint=UserJobDiscoveryService.candidate_evaluation_fingerprint(context),
        evaluation_contract_fingerprint=evaluation_contract_fingerprint_for_runtime(runtime_snapshot_a),
        job_snapshot_json=UserJobDiscoveryService._listing(job).model_dump_json(),
        evaluation_json=opportunity.model_dump_json(),
    ))
    db_session.commit()
    actual_strategy = TailoringStrategy(
        positioning="Evidence-led applied AI engineering.", summary_emphasis="Confirmed service and research work.",
        recommendations=[
            TailoringRecommendation(recommendation_id="requirement-0", requirement_index=0, action="feature", rationale="Confirmed employment evidence.", source_refs=[ApplicationSourceRef(source_type="career_evidence", source_ref=employment_evidence.evidence_id)]),
            TailoringRecommendation(recommendation_id="requirement-1", requirement_index=1, action="feature", rationale="Confirmed project evidence.", source_refs=[ApplicationSourceRef(source_type="career_evidence", source_ref=project_evidence.evidence_id)]),
            TailoringRecommendation(recommendation_id="requirement-2", requirement_index=2, action="feature", rationale="Confirmed Python skill.", source_refs=[ApplicationSourceRef(source_type="skills", source_ref="skills:0")]),
            TailoringRecommendation(recommendation_id="requirement-3", requirement_index=3, action="feature", rationale="Confirmed skill is optional.", source_refs=[ApplicationSourceRef(source_type="skills", source_ref="skills:1")]),
        ], gaps_to_clarify=[], not_to_claim=["Do not claim unconfirmed production ownership."],
    )
    career_evidence_ref = ApplicationSourceRef(source_type="career_evidence", source_ref=employment_evidence.evidence_id)
    project_evidence_ref = ApplicationSourceRef(source_type="career_evidence", source_ref=project_evidence.evidence_id)
    python_skill_ref = ApplicationSourceRef(source_type="skills", source_ref="skills:0")
    actual_draft = CVWritingDraft(
        professional_summary="Engineer building reliable data services and retrieval prototypes.",
        summary_source_refs=[career_evidence_ref, project_evidence_ref, python_skill_ref], key_skills=["Python"],
        role_drafts=[TailoredRoleDraft(employment_index=0, bullets=[TailoredBullet(
            text="Built Python data services.", source_refs=[career_evidence_ref, python_skill_ref, ApplicationSourceRef(source_type="employment", source_ref="employment:0")], priority=80,
        )])],
        project_drafts=[{
            "project_index": 0, "text": "Built and evaluated a retrieval research prototype.",
            "source_refs": [project_evidence_ref, ApplicationSourceRef(source_type="project", source_ref="project:0")], "priority": 75,
        }],
    )
    agent = _Agent(strategy=actual_strategy, draft=actual_draft)
    application_service = ApplicationPreparationService(
        db_session, user_discovery=UserJobDiscoveryService(db_session, runtime_snapshot=runtime_snapshot_a),
        graph_factory=lambda: None, drafting_agent_factory=lambda: None,
        settings=Settings(_env_file=None), runtime_snapshot=runtime_snapshot_a,
    )
    service = ApplicationCVTailoringService(
        db_session, runtime_snapshot=runtime_snapshot_a, application_service=application_service,
        agent_factory=lambda: agent, settings=Settings(_env_file=None),
    )

    preflight = service.preflight(user.id, job.id)
    assert preflight.state == "ready"
    assert preflight.max_requirements == 40 and preflight.max_input_characters == 36_000
    assert agent.strategy_calls == agent.draft_calls == 0
    actual_authority = service._authority(user.id, job.id)
    assert {("career_evidence", employment_evidence.evidence_id), ("career_evidence", project_evidence.evidence_id), ("skills", "skills:0"), ("skills", "skills:1"), ("employment", "employment:0"), ("project", "project:0")} <= set(actual_authority.sources), actual_authority.sources
    service._validate_strategy(actual_strategy, actual_authority)
    started = service.start(user.id, TailoringStartRequest(
        client_request_id="66666666-6666-4666-8666-666666666666", discovered_job_id=job.id,
    ))
    assert started.status == "strategy_review", started.failure_code
    assert started.profile_outline.employment[0].title == "Engineer"
    assert started.profile_outline.projects[0].name == "Agentic research prototype"
    assert agent.strategy_calls == 1
    reviewed = service.review_strategy(user.id, started.id, TailoringStrategyReviewWrite(
        expected_revision=1, accepted=True,
        choices=[
            TailoringStrategyReview(recommendation_id="requirement-0", decision="keep", order=2),
            TailoringStrategyReview(recommendation_id="requirement-1", decision="keep", order=1),
            TailoringStrategyReview(recommendation_id="requirement-2", decision="keep", order=0),
            TailoringStrategyReview(recommendation_id="requirement-3", decision="skip", order=3),
        ],
    ))
    assert reviewed.status == "strategy_accepted"
    generated = service.generate_draft(user.id, started.id, TailoringDraftGenerateRequest(
        client_request_id="67676767-6767-4676-8676-676767676767", expected_strategy_revision=2,
    ))
    assert generated.status == "draft_review", generated.failure_code
    draft_context = agent.draft_contexts[0]
    assert draft_context["allowed_skills"] == ["Python"]
    assert not any(source["source_type"] == "skills" and source["source_ref"] == "skills:1" for source in draft_context["sources"])
    assert {("employment", "employment:0"), ("project", "project:0")} == {
        (item["source_type"], item["source_ref"]) for item in draft_context["structural_attribution_refs"]
    }
    assert all(item["source_type"] != "employment" and item["source_type"] != "project" for item in draft_context["selected_support_refs"])
    choices = draft_context["user_review"]["choices"]
    assert next(item["order"] for item in choices if item["recommendation_id"] == "requirement-2") == 0
    assert next(item["decision"] for item in choices if item["recommendation_id"] == "requirement-3") == "skip"
    rejected_hidden_skill = generated.draft.model_copy(update={"key_skills": ["Python", "HiddenTool"]})
    with pytest.raises(ApplicationCVTailoringError, match="not selected in the accepted strategy"):
        service.review_draft(user.id, started.id, TailoringDraftReviewWrite(expected_revision=1, draft=rejected_hidden_skill))
    approved = service.approve(user.id, started.id, TailoringApproveWrite(expected_strategy_revision=2, expected_draft_revision=1))
    final = ApplicationPreparationReadService(db_session).get(user.id, approved.approved_preparation_id)
    assert final.result.cv.key_skills == ["Python"]
    assert len(final.result.cv.roles) == 1 and final.result.cv.roles[0].bullets[0].text == "Built Python data services."
    assert len(final.result.cv.selected_projects) == 1
    assert final.result.cv.selected_projects[0].name == "Agentic research prototype"
    assert final.result.cv.selected_projects[0].text == "Built and evaluated a retrieval research prototype."
    assert db_session.query(CandidateStructuredProfile).filter_by(user_id=user.id).count() == 1
    assert db_session.query(CandidateCVIngestionDraft).filter_by(user_id=user.id).count() == 0
    assert db_session.query(CVVisibilityAssessment).filter_by(user_id=user.id).count() == 0
    assert db_session.query(ApplicationPreparation).filter_by(user_id=user.id).count() == 1


def test_repeated_start_is_idempotent_and_preserves_uncertain_provider_outcome(db_session, runtime_snapshot_a):
    user = User(id="tailor-uncertain", email="uncertain@example.test", password_hash="safe")
    db_session.add(user); db_session.commit()
    authority = _make_authority(db_session, runtime_snapshot_a, user.id, "uncertain-job")
    agent = _Agent(error=RuntimeError("transport interrupted"))
    service = _service(db_session, runtime_snapshot_a, authority, agent)
    request = TailoringStartRequest(
        client_request_id="33333333-3333-4333-8333-333333333333", discovered_job_id=authority.job.id,
    )
    first = service.start(user.id, request)
    second = service.start(user.id, request)
    assert first.status == "strategy_uncertain"
    assert second.id == first.id
    assert agent.strategy_calls == 1
    assert db_session.query(ApplicationCVTailoringSession).filter_by(user_id=user.id).count() == 1


def test_invalid_strategy_output_is_not_presented_as_completed(db_session, runtime_snapshot_a):
    user = User(id="tailor-invalid", email="invalid@example.test", password_hash="safe")
    db_session.add(user); db_session.commit()
    authority = _make_authority(db_session, runtime_snapshot_a, user.id, "invalid-job")
    agent = _Agent(error=SemanticOutputError("invalid strategy"))
    service = _service(db_session, runtime_snapshot_a, authority, agent)
    result = service.start(user.id, TailoringStartRequest(
        client_request_id="44444444-4444-4444-8444-444444444444", discovered_job_id=authority.job.id,
    ))
    assert result.status == "strategy_invalid"
    assert result.strategy is None
    assert result.strategy_status == "result_invalid"


def test_changed_profile_authority_stales_a_review_before_any_draft_call(db_session, runtime_snapshot_a):
    user = User(id="tailor-stale", email="stale@example.test", password_hash="safe")
    db_session.add(user); db_session.commit()
    authority = _make_authority(db_session, runtime_snapshot_a, user.id, "stale-job")
    agent = _Agent(strategy=_strategy(), draft=_draft())
    service = _service(db_session, runtime_snapshot_a, authority, agent)
    authority_holder = {"current": authority}
    service._authority = lambda _user_id, _job_id, runtime=None: authority_holder["current"]
    started = service.start(user.id, TailoringStartRequest(
        client_request_id="77777777-7777-4777-8777-777777777777", discovered_job_id=authority.job.id,
    ))
    authority_holder["current"] = replace(authority, candidate_fingerprint="changed-profile-fingerprint")
    with pytest.raises(ApplicationCVTailoringError, match="Profile, job evaluation, or AI settings changed"):
        service.review_strategy(user.id, started.id, TailoringStrategyReviewWrite(expected_revision=1, accepted=True))
    current = service.get(user.id, started.id)
    assert current.status == "stale"
    assert current.is_current is False
    assert agent.draft_calls == 0


def test_unsupported_user_draft_edit_returns_validation_error_without_advancing_revision(db_session, runtime_snapshot_a):
    user = User(id="tailor-edit-validation", email="edit-validation@example.test", password_hash="safe")
    db_session.add(user); db_session.commit()
    authority = _make_authority(db_session, runtime_snapshot_a, user.id, "edit-validation-job")
    agent = _Agent(strategy=_strategy(), draft=_draft())
    service = _service(db_session, runtime_snapshot_a, authority, agent)
    started = service.start(user.id, TailoringStartRequest(
        client_request_id="90909090-9090-4090-8090-909090909090", discovered_job_id=authority.job.id,
    ))
    accepted = service.review_strategy(user.id, started.id, TailoringStrategyReviewWrite(expected_revision=1, accepted=True))
    draft = service.generate_draft(user.id, started.id, TailoringDraftGenerateRequest(
        client_request_id="91919191-9191-4191-8191-919191919191", expected_strategy_revision=accepted.strategy_revision,
    ))
    invalid = draft.draft.model_copy(update={"key_skills": ["PyTorch"]})

    with pytest.raises(ApplicationCVTailoringError, match="skill outside confirmed candidate skills") as error:
        service.review_draft(user.id, started.id, TailoringDraftReviewWrite(expected_revision=1, draft=invalid))

    assert error.value.status_code == 422
    unchanged = service.get(user.id, started.id)
    assert unchanged.draft_revision == 1
    assert unchanged.draft.professional_summary == draft.draft.professional_summary


@pytest.mark.parametrize("case,target,claim", [
    ("production", "summary", "Independently owned enterprise production deployments."),
    ("leadership", "bullet", "Led engineering teams through production delivery."),
    ("qualification", "summary", "Holds an active chartered engineering qualification."),
    ("reference-reuse", "summary", "Led production research programs."),
])
def test_user_authored_factual_claims_and_reused_refs_are_unverified_until_acknowledged(db_session, runtime_snapshot_a, case, target, claim):
    user = User(id=f"tailor-claim-{case}", email=f"{case}@example.test", password_hash="safe")
    db_session.add(user); db_session.commit()
    authority = _make_authority(db_session, runtime_snapshot_a, user.id, f"claim-job-{case}")
    agent = _Agent(strategy=_strategy(), draft=_draft())
    service = _service(db_session, runtime_snapshot_a, authority, agent)
    started = service.start(user.id, TailoringStartRequest(
        client_request_id="a9090909-9090-4909-8090-909090909090", discovered_job_id=authority.job.id,
    ))
    accepted = service.review_strategy(user.id, started.id, TailoringStrategyReviewWrite(expected_revision=1, accepted=True))
    draft = service.generate_draft(user.id, started.id, TailoringDraftGenerateRequest(
        client_request_id="b9191919-1919-4191-8191-919191919191", expected_strategy_revision=accepted.strategy_revision,
    ))
    if target == "summary":
        edited = draft.draft.model_copy(update={"professional_summary": claim})
    else:
        role = draft.draft.role_drafts[0]
        bullets = [role.bullets[0].model_copy(update={"text": claim})]
        edited = draft.draft.model_copy(update={"role_drafts": [role.model_copy(update={"bullets": bullets})]})

    saved = service.review_draft(user.id, started.id, TailoringDraftReviewWrite(expected_revision=1, draft=edited))
    assert saved.draft_claim_review_status == "user_edits_unverified"
    assert saved.draft_claim_review_acknowledged is False
    with pytest.raises(ApplicationCVTailoringError, match="Edited factual claims"):
        service.approve(user.id, started.id, TailoringApproveWrite(expected_strategy_revision=2, expected_draft_revision=2))
    assert db_session.query(ApplicationPreparation).filter_by(user_id=user.id).count() == 0

    acknowledged = service.review_draft(user.id, started.id, TailoringDraftReviewWrite(
        expected_revision=2, draft=saved.draft, acknowledge_unverified_claims=True,
    ))
    assert acknowledged.draft_claim_review_status == "user_edits_acknowledged_unverified"
    assert acknowledged.draft_claim_review_acknowledged is True
    history = json.loads(db_session.get(ApplicationCVTailoringSession, started.id).draft_history_json)
    assert history[-1]["factual_review"]["decision"] == "acknowledged_by_user"
    approved = service.approve(user.id, started.id, TailoringApproveWrite(expected_strategy_revision=2, expected_draft_revision=3))
    final = ApplicationPreparationReadService(db_session).get(user.id, approved.approved_preparation_id)
    assert final.result.tailoring_review.draft_claim_status == "user_edits_acknowledged_unverified"
    assert final.result.cv.professional_summary == (claim if target == "summary" else draft.draft.professional_summary)
    if target == "bullet":
        assert final.result.cv.roles[0].bullets[0].text == claim


def test_unsubstantiated_user_metrics_remain_rejected_by_existing_deterministic_guard(db_session, runtime_snapshot_a):
    user = User(id="tailor-unsupported-metric", email="metric@example.test", password_hash="safe")
    db_session.add(user); db_session.commit()
    authority = _make_authority(db_session, runtime_snapshot_a, user.id, "unsupported-metric-job")
    service = _service(db_session, runtime_snapshot_a, authority, _Agent(strategy=_strategy(), draft=_draft()))
    started = service.start(user.id, TailoringStartRequest(
        client_request_id="c9090909-9090-4909-8090-909090909090", discovered_job_id=authority.job.id,
    ))
    accepted = service.review_strategy(user.id, started.id, TailoringStrategyReviewWrite(expected_revision=1, accepted=True))
    draft = service.generate_draft(user.id, started.id, TailoringDraftGenerateRequest(
        client_request_id="d9191919-1919-4191-8191-919191919191", expected_strategy_revision=accepted.strategy_revision,
    ))
    edited = draft.draft.model_copy(update={"professional_summary": "Led a team of 12 engineers."})
    with pytest.raises(ApplicationCVTailoringError, match="unsupported numerical claim"):
        service.review_draft(user.id, started.id, TailoringDraftReviewWrite(
            expected_revision=1, draft=edited, acknowledge_unverified_claims=True,
        ))
    assert service.get(user.id, started.id).draft_revision == 1


def test_positioning_override_requires_recorded_unverified_acknowledgement(db_session, runtime_snapshot_a):
    user = User(id="tailor-positioning-review", email="position@example.test", password_hash="safe")
    db_session.add(user); db_session.commit()
    authority = _make_authority(db_session, runtime_snapshot_a, user.id, "positioning-job")
    service = _service(db_session, runtime_snapshot_a, authority, _Agent(strategy=_strategy(), draft=_draft()))
    started = service.start(user.id, TailoringStartRequest(
        client_request_id="e9090909-9090-4909-8090-909090909090", discovered_job_id=authority.job.id,
    ))
    with pytest.raises(ApplicationCVTailoringError, match="user-written and not verified"):
        service.review_strategy(user.id, started.id, TailoringStrategyReviewWrite(
            expected_revision=1, accepted=True, positioning_override="Owned production deployments and led a large engineering group.",
        ))
    assert service.get(user.id, started.id).strategy_revision == 1
    reviewed = service.review_strategy(user.id, started.id, TailoringStrategyReviewWrite(
        expected_revision=1, accepted=True,
        positioning_override="Owned production deployments and led a large engineering group.",
        acknowledge_unverified_positioning=True,
    ))
    assert reviewed.strategy_positioning_review_status == "user_acknowledged_unverified"
    assert json.loads(db_session.get(ApplicationCVTailoringSession, started.id).strategy_review_json)["positioning_user_decision"] == "acknowledged_unverified_positioning"
    draft = service.generate_draft(user.id, started.id, TailoringDraftGenerateRequest(
        client_request_id="e9292929-2929-4929-8929-292929292929", expected_strategy_revision=2,
    ))
    assert draft.status == "draft_review"
    approved = service.approve(user.id, started.id, TailoringApproveWrite(expected_strategy_revision=2, expected_draft_revision=1))
    final = ApplicationPreparationReadService(db_session).get(user.id, approved.approved_preparation_id)
    assert final.result.tailoring_review.positioning_status == "user_acknowledged_unverified"
    assert final.result.tailoring_review.positioning_user_decision == "acknowledged_unverified_positioning"


def test_authenticated_api_approval_downloads_the_exact_saved_draft_revision(client, db_session, runtime_snapshot_a):
    user = User(id="tailor-api-owner", email="api-owner@example.test", password_hash="safe")
    db_session.add(user); db_session.commit()
    authority = _make_authority(db_session, runtime_snapshot_a, user.id, "tailor-api-job")
    agent = _Agent(strategy=_strategy(), draft=_draft())
    service = _service(db_session, runtime_snapshot_a, authority, agent)
    fastapi_app.dependency_overrides[get_application_cv_tailoring_service] = lambda: service
    headers = {"Authorization": f"Bearer {create_access_token(user.id)}"}
    try:
        started = client.post("/api/v1/applications/tailoring", headers=headers, json={
            "client_request_id": "f1111111-1111-4111-8111-111111111111",
            "discovered_job_id": authority.job.id, "target_pages": 2,
        })
        assert started.status_code == 201, started.text
        tailoring_id = started.json()["id"]
        accepted = client.post(f"/api/v1/applications/tailoring/{tailoring_id}/strategy", headers=headers, json={
            "expected_revision": 1, "accepted": True, "positioning_override": None,
            "choices": [{"recommendation_id": "requirement-0", "decision": "keep", "order": 0}],
        })
        assert accepted.status_code == 200, accepted.text
        generated = client.post(f"/api/v1/applications/tailoring/{tailoring_id}/draft", headers=headers, json={
            "client_request_id": "f2222222-2222-4222-8222-222222222222",
            "expected_strategy_revision": accepted.json()["strategy_revision"],
        })
        assert generated.status_code == 200, generated.text
        edited = generated.json()["draft"]
        edited["professional_summary"] = "Engineer delivering the exact reviewed Python services summary."
        edited["role_drafts"][0]["bullets"][0]["text"] = "Built the exact reviewed Python data services bullet."
        saved = client.patch(f"/api/v1/applications/tailoring/{tailoring_id}/draft", headers=headers, json={
            "expected_revision": generated.json()["draft_revision"], "draft": edited,
            "acknowledge_unverified_claims": True,
        })
        assert saved.status_code == 200, saved.text
        assert saved.json()["draft_claim_review_status"] == "user_edits_acknowledged_unverified"
        approved = client.post(f"/api/v1/applications/tailoring/{tailoring_id}/approve", headers=headers, json={
            "expected_strategy_revision": saved.json()["strategy_revision"],
            "expected_draft_revision": saved.json()["draft_revision"],
        })
        assert approved.status_code == 200, approved.text
        preparation_id = approved.json()["approved_preparation_id"]
        final_read = client.get(f"/api/v1/applications/{preparation_id}", headers=headers)
        assert final_read.status_code == 200
        assert final_read.json()["result"]["cv"]["professional_summary"] == edited["professional_summary"]
        assert final_read.json()["result"]["cv"]["roles"][0]["bullets"][0]["text"] == edited["role_drafts"][0]["bullets"][0]["text"]
        assert final_read.json()["result"]["tailoring_review"]["draft_user_decision"] == "acknowledged_unverified_user_edits"

        download = client.get(f"/api/v1/applications/{preparation_id}/cv.docx", headers=headers)
        assert download.status_code == 200
        assert "attachment" in download.headers["content-disposition"]
        from docx import Document
        document = Document(BytesIO(download.content))
        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        document_text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        assert edited["professional_summary"] in document_text
        assert edited["role_drafts"][0]["bullets"][0]["text"] in document_text
        assert "Engineer building Python services." not in document_text
        pdf = client.get(f"/api/v1/applications/{preparation_id}/cv.pdf", headers=headers)
        assert pdf.status_code == 200
        assert pdf.headers["content-type"].startswith("application/pdf")
        from pypdf import PdfReader
        pdf_text = " ".join((page.extract_text() or "") for page in PdfReader(BytesIO(pdf.content)).pages)
        pdf_text = re.sub(r"\s+", " ", pdf_text)
        assert edited["professional_summary"] in pdf_text
        assert edited["role_drafts"][0]["bullets"][0]["text"] in pdf_text
        assert "Engineer building Python services." not in pdf_text
    finally:
        fastapi_app.dependency_overrides.pop(get_application_cv_tailoring_service, None)
