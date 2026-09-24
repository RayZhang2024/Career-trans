import json
from datetime import datetime, timezone

import pytest
from docx import Document
from pypdf import PdfReader

from app.models.candidate_cv_ingestion import CandidateStructuredProfile
from app.models.candidate_profile import CandidateProfile
from app.models.user import User
from app.models.discovered_job import DiscoveredJob
from app.schemas.application_preparation import (
    ApplicationPrepareRequest, ApplicationQuestionAnswerSet, ApplicationTargetInput, ApplicationSourceRef,
    CVWritingDraft, CoverLetterContent, TailoredBullet, TailoredRoleDraft, TailoredProjectDraft,
)
from app.schemas.candidate import CandidateContext, CareerEvidence, CareerEvidenceProvenance
from app.schemas.cv_ingestion import CandidateCVData, Credential, Education, Employment, Project, Skill
from app.schemas.agentic_discovery import PageContent
from app.schemas.discovery import DiscoveredJobState
from app.core.config import Settings
from app.core.security import create_access_token
from app.api.deps import get_application_preparation_service, validate_semantic_configuration
from app.api import deps as api_deps
from app.providers.llm import SemanticProviderConfigurationError
from app.main import app as fastapi_app
from app.schemas.job import JobProfile, JobRequirement, RequirementCategory, RequirementImportance
from app.schemas.matching import EvidenceRef, EvidenceSourceType, MatchType, RequirementMatch
from app.services.application_document_renderer import ApplicationDocumentRenderer
from app.services.application_preparation_service import ApplicationPreparationService
from app.services.cv_ingestion_service import PersistedCandidateContextLoader
from app.services.llm_runtime import resolve_runtime_snapshot


class _Drafting:
    def __init__(self, *, bad_skill=False, derived_number=False) -> None:
        self.calls = 0; self.bad_skill = bad_skill; self.derived_number = derived_number; self.contexts = []
    def draft_cv(self, context):
        self.calls += 1; self.contexts.append(context)
        text = "Led delivery from 2 hours to 20 minutes" if not self.derived_number else "Improved delivery by 83%"
        return CVWritingDraft(
            professional_summary=text,
            summary_source_refs=[ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1")],
            key_skills=["AWS" if self.bad_skill else "Python"],
            role_drafts=[TailoredRoleDraft(employment_index=0, bullets=[TailoredBullet(text=text, source_refs=[ApplicationSourceRef(source_type=EvidenceSourceType.EMPLOYMENT, source_ref="employment:0"), ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1")], priority=90)])],
            project_drafts=[TailoredProjectDraft(project_index=0, text="Built Python delivery system", source_refs=[ApplicationSourceRef(source_type=EvidenceSourceType.PROJECT, source_ref="project:0")], priority=50)],
        )
    def draft_cover_letter(self, context):
        self.calls += 1
        return CoverLetterContent(body="I led delivery from 2 hours to 20 minutes.", source_refs=[ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1")])
    def draft_answers(self, context):
        return ApplicationQuestionAnswerSet(answers=[{"question": item, "status": "unsupported", "answer": None, "source_refs": []} for item in context["questions"]])


class _Graph:
    def __init__(self): self.calls = 0; self.job_listings = []
    def invoke(self, *, job_text, candidate_context, job_listing=None):
        self.calls += 1
        self.job_listings.append(job_listing)
        requirement = JobRequirement(text="Python delivery", category=RequirementCategory.TECHNICAL, importance=RequirementImportance.ESSENTIAL, source_text="Python delivery")
        return {"job_profile": JobProfile(title="Engineer", requirements=[requirement]), "requirement_matches": [RequirementMatch(requirement_index=0, requirement=requirement, match_type=MatchType.DEMONSTRATED, score=1, evidence_ids=["e1"], evidence_refs=[EvidenceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1", value="Historical value local to this requirement")], reasoning="safe")]}


class _RawTextGraph(_Graph):
    def invoke(self, *, job_text, candidate_context, job_listing=None):
        self.calls += 1
        self.job_listings.append(job_listing)
        requirement = JobRequirement(text="Python delivery", category=RequirementCategory.TECHNICAL, importance=RequirementImportance.ESSENTIAL, source_text="Python delivery")
        profile = JobProfile(title="AI Engineer", company="Brunswick Group", location="London", requirements=[requirement])
        return {"job_profile": profile, "requirement_matches": [RequirementMatch(requirement_index=0, requirement=requirement, match_type=MatchType.DEMONSTRATED, score=1, evidence_ids=["e1"], reasoning="safe")]}


class _Discovery:
    def __init__(self, reusable=None): self.reusable = reusable; self.actionability_calls = 0; self.reuse_calls = 0
    def is_currently_actionable(self, job): self.actionability_calls += 1; return True
    def current_evaluation_for_job(self, user_id, job): self.reuse_calls += 1; return self.reusable


class _Pages:
    def __init__(self, html): self.html = html
    def fetch(self, url): return PageContent(requested_url=url, final_url=url, html=self.html)


def _data():
    return CandidateCVData(
        employment=[Employment(employer="Example", title="Engineer", start_date="2020", end_date="2024", location="London", description="Delivered systems")],
        education=[Education(institution="University", qualification="MSc")],
        credentials=[Credential(name="Certification", credential_type="certification", issued_date="2024")],
        skills=[Skill(name="Python")],
        projects=[Project(name="Canonical Project", description="Built Python delivery system", skills=["Python"])],
    )


def _context():
    return CandidateContext(skills_text="Python", evidence=[CareerEvidence(evidence_id="e1", title="Delivery", text="Led delivery from 2 hours to 20 minutes", skills=["Python"], provenance=[CareerEvidenceProvenance(document_sha256="a" * 64, segment_ids=["s1"])])])


def _service(db_session, monkeypatch, drafting=None, discovery=None, pages=None, settings=None, runtime_snapshot=None):
    user = User(id="u1", email="account@example.test", password_hash="safe")
    profile = CandidateProfile(user_id="u1", display_name="Example Person", preferred_email="old@example.test", location="London")
    db_session.add_all([user, profile, CandidateStructuredProfile(user_id="u1", structured_json=_data().model_dump_json())]); db_session.commit()
    context = _context(); monkeypatch.setattr(PersistedCandidateContextLoader, "load_confirmed", lambda _self, _user: context)
    return ApplicationPreparationService(db_session, graph=_Graph(), drafting_agent=drafting or _Drafting(), user_discovery=discovery or _Discovery(), page_fetcher=pages, settings=settings, runtime_snapshot=runtime_snapshot), context


def test_raw_text_creates_immutable_preparation_and_documents(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    result = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20)))
    assert result.target.source_kind == "job_text" and result.identity.email == "old@example.test" and result.result.actual_pdf_pages <= 2
    renderer = ApplicationDocumentRenderer(); docx = renderer.render_cv_docx(result.identity, result.target, result.result); pdf, pages, status = renderer.render_cv_pdf(result.identity, result.target, result.result)
    assert Document(__import__("io").BytesIO(docx)).paragraphs
    assert PdfReader(__import__("io").BytesIO(pdf)).pages and pages <= 2 and status == "fit"
    docx_text = "\n".join(item.text for item in Document(__import__("io").BytesIO(docx)).paragraphs)
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(__import__("io").BytesIO(pdf)).pages)
    expected_substantive_content = {
        "Example Person", "Led delivery from 2 hours to 20 minutes", "Python",
        "Engineer", "Example", "MSc", "Certification", "Canonical Project",
        "Built Python delivery system",
    }
    normalized_docx = " ".join(docx_text.replace("•", " ").split())
    normalized_pdf = " ".join(pdf_text.replace("•", " ").split())
    assert all(item in normalized_docx and item in normalized_pdf for item in expected_substantive_content)
    assert "Selected Projects" in docx_text and "Selected Projects" in pdf_text
    db_session.query(CandidateProfile).filter_by(user_id="u1").one().preferred_email = "new@example.test"; db_session.commit()
    assert service.get("u1", result.id).identity.email == "old@example.test"


def test_preparation_attribution_distinguishes_fresh_from_reused_analysis_and_is_immutable(db_session, monkeypatch, runtime_snapshot_a):
    from app.schemas.ai_settings import SemanticOperation
    from app.services.llm_runtime import resolve_runtime_snapshot

    fresh_service, _ = _service(db_session, monkeypatch, runtime_snapshot=runtime_snapshot_a)
    fresh = fresh_service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20), include_cover_letter=False))
    expected_fresh = {
        SemanticOperation.JOB_EXTRACTION.value,
        SemanticOperation.REQUIREMENT_MATCHING.value,
        SemanticOperation.CAREER_ALIGNMENT.value,
        SemanticOperation.APPLICATION_DRAFTING.value,
    }
    assert set(fresh.runtime_attribution.operations) == expected_fresh
    assert fresh.runtime_attribution.provider == runtime_snapshot_a.provider
    for operation in expected_fresh:
        resolved = runtime_snapshot_a.operation(operation)
        assert fresh.runtime_attribution.operations[operation].model == resolved.model
        assert fresh.runtime_attribution.operations[operation].reasoning_effort == resolved.reasoning_effort
    from app.models.application_preparation import ApplicationPreparation
    original_json = db_session.get(ApplicationPreparation, fresh.id).runtime_attribution_json

    graph = _Graph()
    graph_state = graph.invoke(job_text="synthetic", candidate_context=_context())
    reusable = type("Reusable", (), {"job_profile": graph_state["job_profile"], "requirement_matches": graph_state["requirement_matches"]})()
    reused_service = fresh_service
    reused_service._user_discovery = _Discovery(reusable)
    reused_service._graph.calls = 0
    reusable_job = _job(); db_session.add(reusable_job); db_session.commit()
    reused = reused_service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(discovered_job_id=reusable_job.id), include_cover_letter=False))
    assert set(reused.runtime_attribution.operations) == {SemanticOperation.APPLICATION_DRAFTING.value}
    assert reused_service._graph.calls == 0

    current_runtime_b = resolve_runtime_snapshot(Settings(default_llm_provider="openai", application_drafting_model="gpt-5.6-sol"))
    fresh_service._runtime_snapshot = current_runtime_b
    assert fresh_service.get("u1", fresh.id).runtime_attribution.model_dump(mode="json") == fresh.runtime_attribution.model_dump(mode="json")
    assert db_session.get(ApplicationPreparation, fresh.id).runtime_attribution_json == original_json
    with pytest.raises(LookupError):
        fresh_service.get("another-user", fresh.id)


def test_legacy_preparation_attribution_projects_unavailable(db_session, monkeypatch):
    from app.models.application_preparation import ApplicationPreparation

    service, _ = _service(db_session, monkeypatch)
    prepared = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20), include_cover_letter=False))
    row = db_session.get(ApplicationPreparation, prepared.id)
    row.runtime_attribution_json = None
    db_session.commit()
    assert service.get("u1", prepared.id).runtime_attribution.status == "legacy_unavailable"


def test_raw_job_text_uses_semantic_metadata_without_synthetic_listing_authority(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    graph = _RawTextGraph()
    service._graph = graph

    result = service.prepare(
        "u1",
        ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="AI Engineer at Brunswick Group. " * 20), include_cover_letter=False),
    )

    assert graph.job_listings == [None]
    assert result.target.source_kind == "job_text"
    assert result.target.title == "AI Engineer"
    assert result.target.company == "Brunswick Group"


def test_unknown_skill_and_derived_number_fail_closed(db_session, monkeypatch):
    from app.models.application_preparation import ApplicationPreparation

    service, _ = _service(db_session, monkeypatch, _Drafting(bad_skill=True))
    with pytest.raises(ValueError, match="skill"):
        service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python role " * 20), include_cover_letter=False))
    assert db_session.query(ApplicationPreparation).count() == 0
    service._drafting_agent = _Drafting(derived_number=True)
    with pytest.raises(ValueError, match="numerical"):
        service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python role " * 20), include_cover_letter=False))


def test_target_must_be_exactly_one_and_unknown_source_rejected(db_session, monkeypatch):
    with pytest.raises(ValueError, match="Exactly one"):
        ApplicationTargetInput(job_text="x" * 100, discovered_job_id="j")
    service, _ = _service(db_session, monkeypatch)
    class Bad(_Drafting):
        def draft_cv(self, context):
            result = super().draft_cv(context); result.summary_source_refs = [ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="missing")]; return result
    service._drafting_agent = Bad()
    with pytest.raises(ValueError, match="unknown or inactive"):
        service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python role " * 20), include_cover_letter=False))


def test_identity_is_checked_before_drafting(db_session, monkeypatch):
    drafting = _Drafting(); service, _ = _service(db_session, monkeypatch, drafting)
    db_session.query(CandidateProfile).filter_by(user_id="u1").one().display_name = None; db_session.commit()
    with pytest.raises(ValueError, match="display name"):
        service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python role " * 20)))
    assert drafting.calls == 0


def test_identity_uses_authenticated_account_email_when_preferred_email_is_absent(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    db_session.query(CandidateProfile).filter_by(user_id="u1").one().preferred_email = None
    db_session.commit()
    result = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20), include_cover_letter=False))
    assert result.identity.email == "account@example.test"


def test_historical_preparation_uses_persisted_candidate_snapshot_after_current_data_changes(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    result = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20)))
    original = service.get("u1", result.id)
    original_result = original.result.model_dump(mode="json")
    db_session.query(CandidateProfile).filter_by(user_id="u1").one().display_name = "Changed Person"
    db_session.query(CandidateStructuredProfile).filter_by(user_id="u1").one().structured_json = CandidateCVData(skills=[Skill(name="Kubernetes")]).model_dump_json()
    db_session.commit()
    historical = service.get("u1", result.id)
    rendered = ApplicationDocumentRenderer().render_cv_docx(historical.identity, historical.target, historical.result)
    assert historical.result.model_dump(mode="json") == original_result
    assert "Example Person" in "\n".join(item.text for item in Document(__import__("io").BytesIO(rendered)).paragraphs)


def test_historical_preparation_keeps_canonical_job_snapshot_after_live_job_changes(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    job = _job(); db_session.add(job); db_session.commit()
    result = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(discovered_job_id=job.id)))
    original_target = result.target.model_dump(mode="json")
    job.title = "Changed title"; job.description = "Changed description"; job.content_hash = "c" * 64
    db_session.commit()
    historical = service.get("u1", result.id)
    cover = ApplicationDocumentRenderer().render_cover_letter_docx(historical.identity, historical.target, historical.result)
    cover_text = "\n".join(item.text for item in Document(__import__("io").BytesIO(cover)).paragraphs)
    assert historical.target.model_dump(mode="json") == original_target
    assert "Dear Company," in cover_text


def test_v2c2_envelope_persists_exact_bounded_drafting_sources_without_changing_public_result(db_session, monkeypatch):
    drafting = _Drafting()
    service, _ = _service(db_session, monkeypatch, drafting=drafting)
    request = ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20), include_cover_letter=False)
    created = service.prepare("u1", request)
    row = db_session.query(__import__("app.models.application_preparation", fromlist=["ApplicationPreparation"]).ApplicationPreparation).filter_by(id=created.id).one()
    persisted = json.loads(row.preparation_result_json)

    assert persisted["persistence_version"] == 2
    assert persisted["evidence_snapshot_status"] == "available"
    assert persisted["result"] == created.result.model_dump(mode="json")
    expected_sources = drafting.contexts[0]["sources"]
    assert persisted["evidence_sources"] == expected_sources
    assert len(expected_sources) <= 18
    assert created.result.model_dump(mode="json").keys() == {"cv", "cover_letter", "answers", "layout_status", "target_pages", "actual_pdf_pages"}


def test_raw_v2c1_result_remains_readable_as_explicit_legacy_unavailable(db_session, monkeypatch):
    from app.models.application_preparation import ApplicationPreparation
    from app.schemas.application_preparation import ApplicationEvidenceSnapshotStatus

    service, _ = _service(db_session, monkeypatch)
    created = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20), include_cover_letter=False))
    row = db_session.query(ApplicationPreparation).filter_by(id=created.id).one()
    original_input_fingerprint = row.preparation_input_fingerprint
    original_contract_fingerprint = row.preparation_contract_fingerprint
    row.preparation_result_json = json.dumps(created.result.model_dump(mode="json"))
    db_session.commit()

    assert service.get("u1", created.id).result == created.result
    assert row.preparation_input_fingerprint == original_input_fingerprint
    assert row.preparation_contract_fingerprint == original_contract_fingerprint
    review = service.get_review("u1", created.id)
    assert review.preparation_id == created.id
    assert review.evidence_snapshot_status == ApplicationEvidenceSnapshotStatus.LEGACY_UNAVAILABLE
    assert review.evidence_sources == []


def test_empty_available_snapshot_is_not_confused_with_legacy_absence(db_session, monkeypatch):
    from app.models.application_preparation import ApplicationPreparation

    service, _ = _service(db_session, monkeypatch)
    created = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20), include_cover_letter=False))
    row = db_session.query(ApplicationPreparation).filter_by(id=created.id).one()
    persisted = json.loads(row.preparation_result_json)
    persisted["evidence_sources"] = []
    row.preparation_result_json = json.dumps(persisted)
    db_session.commit()

    review = service.get_review("u1", created.id)
    assert review.evidence_snapshot_status.value == "available"
    assert review.evidence_sources == []


def test_preparation_input_fingerprint_keeps_v1_semantics_while_contract_bumps(db_session, monkeypatch):
    import app.services.application_preparation_service as preparation_module

    settings = Settings(default_llm_provider="openai", application_drafting_model="gpt-5.6-luna")
    service, _ = _service(db_session, monkeypatch, settings=settings)
    request = ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20), include_cover_letter=False)
    before = service._resolve_target("u1", request, _context()).target
    structured = service._structured("u1")
    catalog = service._bounded_sources(_context(), structured, before.requirement_matches)
    from app.schemas.application_preparation import ApplicationIdentitySnapshot
    identity = ApplicationIdentitySnapshot(display_name="Example Person", email="old@example.test", location="London")
    old_input_payload = {
        "target": before,
        "identity": identity,
        "sources": [
            {"source_type": source_type, "source_ref": source_ref, "text": text}
            for (source_type, source_ref), text in sorted(catalog.items())
        ],
        "structured": structured,
        "options": {"target_pages": request.target_pages, "include_cover_letter": request.include_cover_letter, "application_questions": request.application_questions},
    }
    expected_v1_input = preparation_module._hash(preparation_module._dump(old_input_payload))
    monkeypatch.setattr(preparation_module, "_CONTRACT_VERSION", "application-preparation-v1")
    v1_contract = service._contract_fingerprint()
    monkeypatch.setattr(preparation_module, "_CONTRACT_VERSION", "application-preparation-v2")
    v2_contract = service._contract_fingerprint()

    assert service._input_fingerprint(before, identity, catalog, structured, request) == expected_v1_input
    assert v1_contract != v2_contract


def test_unsupported_questions_are_not_fabricated(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    result = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python role " * 20), include_cover_letter=False, application_questions=["Unsupported question?"]))
    assert result.result.answers[0].status == "unsupported" and result.result.answers[0].answer is None


def test_renderer_compacts_low_priority_bullets_before_overflow():
    from app.schemas.application_preparation import ApplicationIdentitySnapshot, ApplicationPreparationResult, ApplicationTargetSnapshot, TailoredCVContent, TailoredRole
    from app.schemas.job import JobProfile
    role = TailoredRole(employer="Example", title="Engineer", bullets=[TailoredBullet(text="Useful evidence " * 35, source_refs=[ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1")], priority=1) for _ in range(12)])
    result = ApplicationPreparationResult(cv=TailoredCVContent(professional_summary="Supported summary", summary_source_refs=[ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1")], roles=[role]), target_pages=1)
    identity = ApplicationIdentitySnapshot(display_name="Person", email="person@example.test")
    target = ApplicationTargetSnapshot(source_kind="job_text", title="Role", job_profile=JobProfile(title="Role", requirements=[]), job_content_hash="a" * 64)
    compacted, pages, status = ApplicationDocumentRenderer().compact_cv_to_target(identity, target, result)
    assert len(compacted.cv.roles[0].bullets) < len(role.bullets)
    assert status.value in {"fit", "overflow"} and pages >= 1


def test_cover_letter_pdf_is_a_valid_single_page_document(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    preparation = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20)))
    payload, pages, status = ApplicationDocumentRenderer().render_cover_letter_pdf(preparation.identity, preparation.target, preparation.result)
    assert PdfReader(__import__("io").BytesIO(payload)).pages
    assert pages <= 1 and status.value == "fit"


def _job() -> DiscoveredJob:
    now = datetime.now(timezone.utc)
    return DiscoveredJob(identity_key="provider:application", source="greenhouse", source_token="board", external_id="application", title="Canonical Job", company="Company", location="London", url="https://jobs.example/application", description="Python delivery role " * 30, detail_authority="verified_employer_detail", verification_status="verified", content_hash="b" * 64, state=DiscoveredJobState.UNCHANGED.value, first_seen_at=now, last_seen_at=now, last_changed_at=now)


def test_canonical_target_uses_shared_authority_and_graph_when_not_reusable(db_session, monkeypatch):
    discovery = _Discovery(); service, _ = _service(db_session, monkeypatch, discovery=discovery)
    job = _job(); db_session.add(job); db_session.commit()
    graph = service._graph
    result = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(discovered_job_id=job.id), include_cover_letter=False))
    assert result.target.canonical_discovered_job_id == job.id
    assert discovery.actionability_calls == 1 and discovery.reuse_calls == 1 and graph.calls == 1
    assert graph.job_listings[0] is not None
    assert graph.job_listings[0].title == "Canonical Job"


def test_reusable_canonical_analysis_bypasses_graph(db_session, monkeypatch):
    graph = _Graph(); state = graph.invoke(job_text="x", candidate_context=_context())
    reusable = type("Reusable", (), {"job_profile": state["job_profile"], "requirement_matches": state["requirement_matches"]})()
    discovery = _Discovery(reusable); service, _ = _service(db_session, monkeypatch, discovery=discovery); service._graph = graph; graph.calls = 0
    job = _job(); db_session.add(job); db_session.commit()
    service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(discovered_job_id=job.id), include_cover_letter=False))
    assert discovery.reuse_calls == 1 and graph.calls == 0


def test_reusable_target_resolution_does_not_construct_semantic_clients(db_session, monkeypatch):
    service, context = _service(db_session, monkeypatch)
    graph = _Graph()
    state = graph.invoke(job_text="Python delivery role", candidate_context=context)
    reusable = type("Reusable", (), {"job_profile": state["job_profile"], "requirement_matches": state["requirement_matches"]})()
    settings = Settings(default_llm_provider="openai")
    runtime = resolve_runtime_snapshot(settings)
    monkeypatch.setattr(api_deps, "get_settings", lambda: settings)

    def provider_construction_is_not_needed(*args, **kwargs):
        pytest.fail("reusable target lookup must not construct a semantic provider client")

    monkeypatch.setattr(api_deps, "get_semantic_response_client", provider_construction_is_not_needed)
    prepared_service = api_deps.get_application_preparation_service(db_session, runtime)
    prepared_service._user_discovery = _Discovery(reusable)
    job = _job(); db_session.add(job); db_session.commit()

    target = prepared_service._resolve_target(
        "u1",
        ApplicationPrepareRequest(target=ApplicationTargetInput(discovered_job_id=job.id)),
        context,
    ).target

    assert target.canonical_discovered_job_id == job.id
    assert prepared_service._graph is None and prepared_service._drafting_agent is None


def test_url_requires_jobposting_metadata_and_stays_local(db_session, monkeypatch):
    description = "Python delivery responsibilities and requirements. " * 20
    html = '<script type="application/ld+json">' + json.dumps({"@type": "JobPosting", "title": "URL Engineer", "description": description}) + "</script>"
    service, _ = _service(db_session, monkeypatch, pages=_Pages(html))
    before = db_session.query(DiscoveredJob).count()
    result = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_url="https://jobs.example/url"), include_cover_letter=False))
    assert result.target.source_kind == "job_url" and db_session.query(DiscoveredJob).count() == before
    assert service._graph.job_listings[0] is not None
    assert service._graph.job_listings[0].title == "URL Engineer"


def test_url_without_plausible_vacancy_fails_safely(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch, pages=_Pages("plain page " * 100))
    from app.schemas.application_preparation import ApplicationInsufficientDetailError
    with pytest.raises(ApplicationInsufficientDetailError):
        service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_url="https://jobs.example/no-job"), include_cover_letter=False))


def test_essential_requirement_evidence_precedes_structural_anchor_budget(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    data = _data()
    data.employment.extend(
        Employment(employer=f"Employer {index}", title="Engineer", start_date=str(2000 + index), description="Delivered systems")
        for index in range(30)
    )
    context = _context().model_copy(update={"evidence": _context().evidence + [CareerEvidence(evidence_id="essential-evidence", title="Essential", text="Essential Python delivery evidence", skills=["Python"], provenance=[CareerEvidenceProvenance(document_sha256="b" * 64, segment_ids=["s2"])]) ]})
    requirement = JobRequirement(text="Python delivery", category=RequirementCategory.TECHNICAL, importance=RequirementImportance.ESSENTIAL, source_text="Python delivery")
    match = RequirementMatch(requirement_index=0, requirement=requirement, match_type=MatchType.DEMONSTRATED, score=1, evidence_ids=["essential-evidence"], reasoning="safe")
    catalog = service._bounded_sources(context, data, [match])
    assert (EvidenceSourceType.CAREER_EVIDENCE.value, "essential-evidence") in catalog
    assert len(catalog) <= 18
    target = service._resolve_target("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python role " * 20)), context).target
    draft_context = service._draft_context(target, catalog, data)
    assert draft_context["employment"]
    assert all((EvidenceSourceType.EMPLOYMENT.value, item["required_anchor_ref"]["source_ref"]) in catalog for item in draft_context["employment"])


def test_role_attribution_project_and_prose_skill_boundaries(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    data = _data()
    data.employment.append(Employment(employer="Second", title="Architect", start_date="2018", end_date="2020", description="Designed systems"))
    catalog = service._bounded_sources(_context(), data, _Graph().invoke(job_text="x", candidate_context=_context())["requirement_matches"])
    good = _Drafting().draft_cv({})
    good.project_drafts = [TailoredProjectDraft(project_index=0, text="Built Python delivery system", source_refs=[ApplicationSourceRef(source_type=EvidenceSourceType.PROJECT, source_ref="project:0")], priority=50)]
    assert service._materialize_cv(good, data, catalog).selected_projects[0].name == "Canonical Project"
    bad_role = good.model_copy(deep=True)
    bad_role.role_drafts.append(TailoredRoleDraft(employment_index=1, bullets=[TailoredBullet(text="Designed systems", source_refs=[ApplicationSourceRef(source_type=EvidenceSourceType.EMPLOYMENT, source_ref="employment:0")], priority=50)]))
    with pytest.raises(ValueError, match="canonical employment attribution"):
        service._materialize_cv(bad_role, data, catalog)
    duplicate_role = good.model_copy(deep=True)
    duplicate_role.role_drafts.append(
        TailoredRoleDraft(employment_index=0, bullets=[])
    )
    with pytest.raises(ValueError, match="duplicate employment drafts"):
        service._materialize_cv(duplicate_role, data, catalog)
    with pytest.raises(ValueError, match="unsupported skill"):
        service._validate_text_and_refs("Built AWS platform", [ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1")], catalog)


@pytest.mark.parametrize("surface", ["summary", "bullet", "cover letter", "drafted answer"])
def test_unsupported_technology_is_rejected_for_every_drafted_prose_surface(db_session, monkeypatch, surface):
    service, _ = _service(db_session, monkeypatch)
    catalog = service._bounded_sources(_context(), _data(), _Graph().invoke(job_text="x", candidate_context=_context())["requirement_matches"])
    with pytest.raises(ValueError, match="unsupported skill or technology"):
        service._validate_text_and_refs(
            "Delivered a Kubernetes platform", [ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1")], catalog,
        )


def test_typed_provenance_refs_are_accepted_only_from_the_bounded_catalog(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    catalog = {
        (EvidenceSourceType.CAREER_EVIDENCE.value, "e1"): "Python delivery",
        (EvidenceSourceType.EMPLOYMENT.value, "employment:0"): "Engineer at Example",
        (EvidenceSourceType.EDUCATION.value, "education:0"): "MSc at University",
        (EvidenceSourceType.CANDIDATE_PROFILE.value, "career_profile"): "Applied AI engineer",
        (EvidenceSourceType.CANDIDATE_ELIGIBILITY.value, "locations"): "London",
    }
    for ref in (
        ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1"),
        ApplicationSourceRef(source_type=EvidenceSourceType.EMPLOYMENT, source_ref="employment:0"),
        ApplicationSourceRef(source_type=EvidenceSourceType.EDUCATION, source_ref="education:0"),
        ApplicationSourceRef(source_type=EvidenceSourceType.CANDIDATE_PROFILE, source_ref="career_profile"),
        ApplicationSourceRef(source_type=EvidenceSourceType.CANDIDATE_ELIGIBILITY, source_ref="locations"),
    ):
        service._validate_text_and_refs("Supported statement", [ref], catalog)
    with pytest.raises(ValueError, match="unknown or inactive"):
        service._validate_text_and_refs("Unsupported statement", [ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="inactive")], catalog)


@pytest.mark.parametrize(
    ("surface", "text", "source_type", "source_ref"),
    [
        ("cover letter", "I delivered Python systems [db217c96-c103-4eb1-9061-be43457422e9].", EvidenceSourceType.CAREER_EVIDENCE, "db217c96-c103-4eb1-9061-be43457422e9"),
        ("summary", "db217c96-c103-4eb1-9061-be43457422e9 delivered Python systems.", EvidenceSourceType.CAREER_EVIDENCE, "db217c96-c103-4eb1-9061-be43457422e9"),
        ("role bullet", "employment:0 delivered Python systems.", EvidenceSourceType.EMPLOYMENT, "employment:0"),
        ("selected project", "project:0 delivered Python systems.", EvidenceSourceType.PROJECT, "project:0"),
        ("drafted answer", "db217c96-c103-4eb1-9061-be43457422e9 delivered Python systems.", EvidenceSourceType.CAREER_EVIDENCE, "db217c96-c103-4eb1-9061-be43457422e9"),
    ],
)
def test_internal_provenance_identifiers_are_rejected_before_numeric_validation(db_session, monkeypatch, surface, text, source_type, source_ref):
    service, _ = _service(db_session, monkeypatch)
    catalog = {
        (EvidenceSourceType.CAREER_EVIDENCE.value, "db217c96-c103-4eb1-9061-be43457422e9"): "Delivered Python systems.",
        (EvidenceSourceType.EMPLOYMENT.value, "employment:0"): "Engineer at Example.",
        (EvidenceSourceType.PROJECT.value, "project:0"): "Project delivery.",
    }
    with pytest.raises(ValueError, match="internal provenance identifier"):
        service._validate_text_and_refs(text, [ApplicationSourceRef(source_type=source_type, source_ref=source_ref)], catalog)


def test_unknown_uuid_in_prose_is_rejected_as_identifier_hygiene(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    catalog = {(EvidenceSourceType.CAREER_EVIDENCE.value, "evidence"): "Delivered Python systems."}
    with pytest.raises(ValueError, match="internal provenance identifier"):
        service._validate_text_and_refs("Delivered systems 123e4567-e89b-12d3-a456-426614174000.", [ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="evidence")], catalog)


def test_natural_language_source_refs_are_allowed_in_ordinary_prose(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    catalog = {(EvidenceSourceType.CANDIDATE_ELIGIBILITY.value, "locations"): "London"}
    ref = ApplicationSourceRef(source_type=EvidenceSourceType.CANDIDATE_ELIGIBILITY, source_ref="locations")
    service._validate_text_and_refs("Worked across multiple customer locations.", [ref], catalog)
    service._validate_text_and_refs("The role covers locations across London.", [ref], catalog)


@pytest.mark.parametrize(
    "source_type,source_ref",
    [
        (EvidenceSourceType.EMPLOYMENT, "employment:0"),
        (EvidenceSourceType.PROJECT, "project:0"),
        (EvidenceSourceType.EDUCATION, "education:0"),
        (EvidenceSourceType.CREDENTIAL, "credential:0"),
        (EvidenceSourceType.CANDIDATE_PROFILE, "career_strategy"),
        (EvidenceSourceType.CANDIDATE_ELIGIBILITY, "work_authorisation"),
    ],
)
def test_structural_source_refs_remain_rejected_in_prose(db_session, monkeypatch, source_type, source_ref):
    service, _ = _service(db_session, monkeypatch)
    catalog = {(source_type.value, source_ref): "Supported statement."}
    with pytest.raises(ValueError, match="internal provenance identifier"):
        service._validate_text_and_refs(f"Delivered {source_ref} work.", [ApplicationSourceRef(source_type=source_type, source_ref=source_ref)], catalog)


@pytest.mark.parametrize("text", ["[locations]", "`locations`", "(source_ref: locations)"])
def test_citation_like_natural_source_refs_are_rejected(db_session, monkeypatch, text):
    service, _ = _service(db_session, monkeypatch)
    catalog = {(EvidenceSourceType.CANDIDATE_ELIGIBILITY.value, "locations"): "London"}
    ref = ApplicationSourceRef(source_type=EvidenceSourceType.CANDIDATE_ELIGIBILITY, source_ref="locations")
    with pytest.raises(ValueError, match="internal provenance identifier"):
        service._validate_text_and_refs(text, [ref], catalog)


def test_literal_employment_dates_support_only_literal_date_claims(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    data = _data().model_copy(update={"employment": [Employment(employer="Example", title="Engineer", start_date="2023", end_date="Present", location="London", description="Delivered systems")]})
    catalog = service._bounded_sources(_context(), data, _Graph().invoke(job_text="x", candidate_context=_context())["requirement_matches"])
    employment_key = (EvidenceSourceType.EMPLOYMENT.value, "employment:0")
    assert "2023" in catalog[employment_key]
    ref = ApplicationSourceRef(source_type=EvidenceSourceType.EMPLOYMENT, source_ref="employment:0")
    service._validate_text_and_refs("Engineer at Example since 2023.", [ref], catalog)
    with pytest.raises(ValueError, match="unsupported numerical"):
        service._validate_text_and_refs("Engineer at Example for 3 years.", [ref], catalog)


def test_supported_literal_numeric_claim_passes_but_derived_percentage_fails(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    catalog = {(EvidenceSourceType.CAREER_EVIDENCE.value, "e1"): "Reduced document processing time from 90 minutes to 20 minutes."}
    ref = ApplicationSourceRef(source_type=EvidenceSourceType.CAREER_EVIDENCE, source_ref="e1")
    service._validate_text_and_refs("Reduced document processing time from 90 minutes to 20 minutes.", [ref], catalog)
    with pytest.raises(ValueError, match="unsupported numerical"):
        service._validate_text_and_refs("Reduced processing time by 78%.", [ref], catalog)


def test_actual_contract_and_input_fingerprints_change_only_for_contract_or_options(db_session, monkeypatch):
    settings = Settings(default_llm_provider="openai", application_drafting_model="gpt-5.6-luna")
    service, _ = _service(db_session, monkeypatch, settings=settings)
    first = service._contract_fingerprint(); assert first == service._contract_fingerprint()
    service._settings = Settings(default_llm_provider="ollama", application_drafting_model="gpt-5.6-luna")
    service._runtime_snapshot = resolve_runtime_snapshot(service._settings)
    provider_changed = service._contract_fingerprint()
    service._settings = Settings(default_llm_provider="openai", application_drafting_model="gpt-5.6-sol")
    service._runtime_snapshot = resolve_runtime_snapshot(service._settings)
    model_changed = service._contract_fingerprint()
    assert first != provider_changed and first != model_changed
    service._settings = settings
    service._runtime_snapshot = resolve_runtime_snapshot(settings)
    target = service._resolve_target("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python role " * 20)), _context()).target
    # Request options are part of the preparation input identity, not evaluation reuse.
    source = service._bounded_sources(_context(), _data(), target.requirement_matches)
    from app.schemas.application_preparation import ApplicationIdentitySnapshot
    person = ApplicationIdentitySnapshot(display_name="Person", email="person@example.test")
    a = service._input_fingerprint(target, person, source, _data(), ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python role " * 20)))
    b = service._input_fingerprint(target, person, source, _data(), ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python role " * 20), target_pages=1))
    assert a != b


def test_contract_fingerprint_includes_renderer_version(db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch, settings=Settings(default_llm_provider="openai", application_drafting_model="gpt-5.6-luna"))
    first = service._contract_fingerprint(); service._renderer_template_version = "ats_standard-v2"
    assert first != service._contract_fingerprint()


def test_application_drafting_model_is_a_required_semantic_configuration_value():
    with pytest.raises(SemanticProviderConfigurationError, match="APPLICATION_DRAFTING_MODEL"):
        validate_semantic_configuration(Settings(application_drafting_model=" "))


def test_authenticated_api_owner_and_cross_user_download_isolation(client, db_session, monkeypatch):
    service, _ = _service(db_session, monkeypatch)
    other = User(id="u2", email="other@example.test", password_hash="safe"); db_session.add(other); db_session.commit()
    fastapi_app.dependency_overrides[get_application_preparation_service] = lambda: service
    try:
        owner = {"Authorization": f"Bearer {create_access_token('u1')}"}
        other_headers = {"Authorization": f"Bearer {create_access_token('u2')}"}
        payload = {"target": {"job_text": "Python delivery role " * 20}, "include_cover_letter": True}
        created = client.post("/api/v1/applications/prepare", json=payload, headers=owner)
        assert created.status_code == 201
        preparation_id = created.json()["id"]
        assert client.get("/api/v1/applications", headers=owner).status_code == 200
        review = client.get(f"/api/v1/applications/{preparation_id}/review", headers=owner)
        assert review.status_code == 200
        assert review.json()["preparation_id"] == preparation_id
        assert review.json()["evidence_snapshot_status"] == "available"
        assert review.json()["evidence_sources"]
        assert client.get(f"/api/v1/applications/{preparation_id}/review", headers=other_headers).status_code == 404
        assert client.get("/api/v1/applications/not-owned/review", headers=owner).status_code == 404
        from app.models.application_preparation import ApplicationPreparation
        persisted = db_session.query(ApplicationPreparation).filter_by(id=preparation_id).one()
        persisted.preparation_result_json = json.dumps(created.json()["result"])
        db_session.commit()
        legacy_review = client.get(f"/api/v1/applications/{preparation_id}/review", headers=owner)
        assert legacy_review.status_code == 200
        assert legacy_review.json()["evidence_snapshot_status"] == "legacy_unavailable"
        assert legacy_review.json()["evidence_sources"] == []
        for suffix, media_type in (("cv.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), ("cv.pdf", "application/pdf"), ("cover-letter.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), ("cover-letter.pdf", "application/pdf")):
            download = client.get(f"/api/v1/applications/{preparation_id}/{suffix}", headers=owner)
            assert download.status_code == 200
            assert download.headers["content-type"].startswith(media_type)
            assert "attachment" in download.headers["content-disposition"]
        assert client.get(f"/api/v1/applications/{preparation_id}", headers=other_headers).status_code == 404
        for suffix in ("cv.docx", "cv.pdf", "cover-letter.docx", "cover-letter.pdf"):
            assert client.get(f"/api/v1/applications/{preparation_id}/{suffix}", headers=other_headers).status_code == 404
        injected = {**payload, "candidate_context": {"evidence": []}}
        assert client.post("/api/v1/applications/prepare", json=injected, headers=owner).status_code == 422
    finally:
        pass
        fastapi_app.dependency_overrides.pop(get_application_preparation_service, None)


def test_application_history_detail_and_downloads_are_provider_free(client, db_session, monkeypatch):
    """Persisted application reads do not construct ranking or semantic clients."""
    service, _ = _service(db_session, monkeypatch)
    historical_context = _context().model_copy(update={"evidence": [CareerEvidence(evidence_id="e1", title="Private fixture", text="SNAPSHOT_ONLY_9281 Python delivery from 2 hours to 20 minutes", skills=["Python"], provenance=[CareerEvidenceProvenance(document_sha256="a" * 64, segment_ids=["s1"])])]})
    monkeypatch.setattr(PersistedCandidateContextLoader, "load_confirmed", lambda _self, _user: historical_context)
    with_cover = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20), include_cover_letter=True))
    without_cover = service.prepare("u1", ApplicationPrepareRequest(target=ApplicationTargetInput(job_text="Python delivery role " * 20), include_cover_letter=False))
    db_session.add(User(id="u2", email="other@example.test", password_hash="safe")); db_session.commit()

    constructed = []
    def provider_construction_would_503(*args, **kwargs):
        constructed.append((args, kwargs))
        from fastapi import HTTPException
        raise HTTPException(status_code=503, detail="provider unavailable")

    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setattr(api_deps, "get_semantic_response_client", provider_construction_would_503)
    # Semantic service construction is request-scoped, so a prior direct
    # builder call cannot mask this provider dependency chain.
    owner = {"Authorization": f"Bearer {create_access_token('u1')}"}
    other = {"Authorization": f"Bearer {create_access_token('u2')}"}
    try:
        listing = client.get("/api/v1/applications", headers=owner)
        assert listing.status_code == 200 and {row["id"] for row in listing.json()} == {with_cover.id, without_cover.id}
        assert all("evidence_sources" not in row["result"] and "evidence_snapshot_status" not in row["result"] for row in listing.json())
        assert "SNAPSHOT_ONLY_9281" not in listing.text
        detail = client.get(f"/api/v1/applications/{with_cover.id}", headers=owner)
        assert detail.status_code == 200 and detail.json()["identity"]["display_name"] == "Example Person"
        assert "evidence_sources" not in detail.json()["result"]
        assert "SNAPSHOT_ONLY_9281" not in detail.text
        assert client.get(f"/api/v1/applications/{with_cover.id}", headers=other).status_code == 404
        review = client.get(f"/api/v1/applications/{with_cover.id}/review", headers=owner)
        assert review.status_code == 200 and review.json()["evidence_snapshot_status"] == "available"
        assert review.json()["preparation_id"] == with_cover.id
        assert any("SNAPSHOT_ONLY_9281" in source["text"] for source in review.json()["evidence_sources"])
        assert client.get(f"/api/v1/applications/{with_cover.id}/review", headers=other).status_code == 404
        assert client.get("/api/v1/applications/missing/review", headers=owner).status_code == 404

        for preparation in (with_cover, without_cover):
            for suffix in ("cv.docx", "cv.pdf"):
                response = client.get(f"/api/v1/applications/{preparation.id}/{suffix}", headers=owner)
                assert response.status_code == 200 and "attachment" in response.headers["content-disposition"]
                assert client.get(f"/api/v1/applications/{preparation.id}/{suffix}", headers=other).status_code == 404
        for suffix in ("cover-letter.docx", "cover-letter.pdf"):
            assert client.get(f"/api/v1/applications/{with_cover.id}/{suffix}", headers=owner).status_code == 200
            assert client.get(f"/api/v1/applications/{with_cover.id}/{suffix}", headers=other).status_code == 404
            assert client.get(f"/api/v1/applications/{without_cover.id}/{suffix}", headers=owner).status_code == 404
        assert constructed == []

        # Creation remains intentionally attached to the semantic application
        # service and therefore fails safely when semantic construction fails.
        created = client.post("/api/v1/applications/prepare", headers=owner, json={"target": {"job_text": "Python delivery role " * 20}, "include_cover_letter": False})
        assert created.status_code == 503
        assert constructed
    finally:
        pass
