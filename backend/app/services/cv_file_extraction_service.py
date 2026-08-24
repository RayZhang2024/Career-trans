import hashlib
import io
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.schemas.cv_ingestion import CVDocumentProvenance, ExtractedCVDocument, ExtractedCVSegment


class CVFileExtractionError(ValueError):
    pass


class CVFileExtractionService:
    _allowed = {".pdf", ".docx", ".md", ".markdown", ".json"}
    _max_bytes = 5 * 1024 * 1024

    def extract(self, *, filename: str, content_type: str | None, content: bytes) -> ExtractedCVDocument:
        suffix = Path(filename).suffix.casefold()
        if suffix not in self._allowed:
            raise CVFileExtractionError("Unsupported CV file type. Use PDF, DOCX, Markdown, or JSON.")
        if not content or len(content) > self._max_bytes:
            raise CVFileExtractionError("CV file is empty or exceeds the 5 MB limit.")
        digest = hashlib.sha256(content).hexdigest()
        if suffix == ".pdf":
            try:
                texts = [page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages]
            except Exception as exc:
                raise CVFileExtractionError("PDF could not be read as a text CV.") from exc
        elif suffix == ".docx":
            texts = ["\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs)]
        else:
            try:
                texts = [content.decode("utf-8")]
            except UnicodeDecodeError as exc:
                raise CVFileExtractionError("CV text files must be UTF-8 encoded.") from exc
        segments = [
            ExtractedCVSegment(segment_id=f"{digest}:{index + 1}", text=text.strip(), page_number=index + 1 if suffix == ".pdf" else None)
            for index, text in enumerate(texts)
            if text.strip()
        ]
        if not segments:
            raise CVFileExtractionError("No extractable text was found in the CV file.")
        return ExtractedCVDocument(
            provenance=CVDocumentProvenance(filename=filename, media_type=content_type or "application/octet-stream", document_sha256=digest, segment_ids=[segment.segment_id for segment in segments]),
            segments=segments,
        )

    @staticmethod
    def parsed_json(document: ExtractedCVDocument) -> dict[str, object] | None:
        if not document.provenance.filename.casefold().endswith(".json"):
            return None
        try:
            value = json.loads("\n".join(segment.text for segment in document.segments))
        except json.JSONDecodeError as exc:
            raise CVFileExtractionError("CV JSON must be valid JSON.") from exc
        if not isinstance(value, dict):
            raise CVFileExtractionError("CV JSON must be an object matching the candidate CV schema.")
        return value
