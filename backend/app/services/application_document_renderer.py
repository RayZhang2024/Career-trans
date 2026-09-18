"""Deterministic ATS-standard DOCX/PDF rendering from immutable snapshots."""

from io import BytesIO
import re

from docx import Document
from docx.shared import Cm, Pt
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from app.schemas.application_preparation import (
    ApplicationIdentitySnapshot, ApplicationLayoutStatus, ApplicationPreparationResult,
    ApplicationTargetSnapshot,
)


class ApplicationDocumentRenderer:
    """One substantive document model rendered two ways; no mutable profile reads."""

    BODY_SIZE = 10.5
    MIN_BODY_SIZE = 9.5

    def render_cv_docx(self, identity: ApplicationIdentitySnapshot, target: ApplicationTargetSnapshot, result: ApplicationPreparationResult) -> bytes:
        doc = Document(); section = doc.sections[0]
        section.top_margin = section.bottom_margin = Cm(1.8); section.left_margin = section.right_margin = Cm(1.8)
        styles = doc.styles; styles["Normal"].font.name = "Arial"; styles["Normal"].font.size = Pt(self.BODY_SIZE)
        doc.add_heading(identity.display_name, 0)
        doc.add_paragraph(self._contact(identity))
        self._docx_heading(doc, "Professional Summary"); doc.add_paragraph(result.cv.professional_summary)
        if result.cv.key_skills: self._docx_heading(doc, "Key Skills"); doc.add_paragraph(" • ".join(result.cv.key_skills))
        self._docx_heading(doc, "Experience")
        for role in result.cv.roles:
            doc.add_paragraph(f"{role.title} — {role.employer}").runs[0].bold = True
            doc.add_paragraph(" | ".join(value for value in [self._dates(role.start_date, role.end_date), role.location] if value))
            for bullet in role.bullets: doc.add_paragraph(bullet.text, style="List Bullet")
        if result.cv.education: self._docx_heading(doc, "Education"); [doc.add_paragraph(item) for item in result.cv.education]
        if result.cv.credentials: self._docx_heading(doc, "Credentials"); [doc.add_paragraph(item) for item in result.cv.credentials]
        output = BytesIO(); doc.save(output); return output.getvalue()

    def render_cover_letter_docx(self, identity: ApplicationIdentitySnapshot, target: ApplicationTargetSnapshot, result: ApplicationPreparationResult) -> bytes:
        doc = Document(); doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(self.BODY_SIZE)
        doc.add_paragraph(identity.display_name); doc.add_paragraph(self._contact(identity)); doc.add_paragraph(f"Dear {target.company or 'Hiring Team'},")
        doc.add_paragraph(result.cover_letter.body if result.cover_letter else "")
        doc.add_paragraph("Sincerely,"); doc.add_paragraph(identity.display_name)
        output = BytesIO(); doc.save(output); return output.getvalue()

    def render_cv_pdf(self, identity: ApplicationIdentitySnapshot, target: ApplicationTargetSnapshot, result: ApplicationPreparationResult) -> tuple[bytes, int, ApplicationLayoutStatus]:
        story = self._cv_story(identity, result)
        return self._pdf(story, result.target_pages)

    def compact_cv_to_target(
        self, identity: ApplicationIdentitySnapshot, target: ApplicationTargetSnapshot, result: ApplicationPreparationResult,
    ) -> tuple[ApplicationPreparationResult, int, ApplicationLayoutStatus]:
        """Drop only optional lower-priority bullets before reporting overflow.

        This is deliberately deterministic and has no semantic/model retry.  It
        never removes chronology, education, credentials, or changes font size.
        """
        current = result
        _, pages, status = self.render_cv_pdf(identity, target, current)
        while status == ApplicationLayoutStatus.OVERFLOW:
            candidates = [
                (bullet.priority, -role_index, bullet_index)
                for role_index, role in enumerate(current.cv.roles)
                for bullet_index, bullet in enumerate(role.bullets)
                if bullet.priority < 100
            ]
            if not candidates:
                break
            _, negative_role_index, bullet_index = min(candidates)
            role_index = -negative_role_index
            roles = list(current.cv.roles); role = roles[role_index]
            bullets = list(role.bullets); bullets.pop(bullet_index)
            roles[role_index] = role.model_copy(update={"bullets": bullets})
            current = current.model_copy(update={"cv": current.cv.model_copy(update={"roles": roles})})
            _, pages, status = self.render_cv_pdf(identity, target, current)
        return current, pages, status

    def render_cover_letter_pdf(self, identity: ApplicationIdentitySnapshot, target: ApplicationTargetSnapshot, result: ApplicationPreparationResult) -> tuple[bytes, int, ApplicationLayoutStatus]:
        normal, title = self._styles()
        story = [Paragraph(identity.display_name, title), Paragraph(self._contact(identity), normal), Spacer(1, 10), Paragraph(f"Dear {target.company or 'Hiring Team'},", normal), Spacer(1, 8), Paragraph(_paragraph(result.cover_letter.body if result.cover_letter else ""), normal), Spacer(1, 8), Paragraph("Sincerely,<br/>" + identity.display_name, normal)]
        return self._pdf(story, 1)

    def _cv_story(self, identity: ApplicationIdentitySnapshot, result: ApplicationPreparationResult):
        normal, title = self._styles(); heading = ParagraphStyle("heading", parent=normal, fontName="Helvetica-Bold", spaceBefore=8, spaceAfter=3)
        story = [Paragraph(identity.display_name, title), Paragraph(self._contact(identity), normal), Paragraph("Professional Summary", heading), Paragraph(_paragraph(result.cv.professional_summary), normal)]
        if result.cv.key_skills: story.extend([Paragraph("Key Skills", heading), Paragraph(" • ".join(result.cv.key_skills), normal)])
        story.append(Paragraph("Experience", heading))
        for role in result.cv.roles:
            story.append(Paragraph(f"<b>{role.title} — {role.employer}</b>", normal)); meta = " | ".join(value for value in [self._dates(role.start_date, role.end_date), role.location] if value)
            if meta: story.append(Paragraph(meta, normal))
            for bullet in role.bullets: story.append(Paragraph("• " + _paragraph(bullet.text), normal))
        if result.cv.education: story.append(Paragraph("Education", heading)); story.extend(Paragraph(_paragraph(item), normal) for item in result.cv.education)
        if result.cv.credentials: story.append(Paragraph("Credentials", heading)); story.extend(Paragraph(_paragraph(item), normal) for item in result.cv.credentials)
        return story

    def _pdf(self, story, target_pages: int) -> tuple[bytes, int, ApplicationLayoutStatus]:
        output = BytesIO(); doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=1.8*cm, leftMargin=1.8*cm, topMargin=1.6*cm, bottomMargin=1.6*cm)
        doc.build(story); payload = output.getvalue(); pages = len(PdfReader(BytesIO(payload)).pages)
        return payload, pages, ApplicationLayoutStatus.FIT if pages <= target_pages else ApplicationLayoutStatus.OVERFLOW

    def _styles(self):
        base = getSampleStyleSheet(); normal = ParagraphStyle("career_body", parent=base["Normal"], fontName="Helvetica", fontSize=self.BODY_SIZE, leading=13)
        return normal, ParagraphStyle("career_name", parent=normal, fontName="Helvetica-Bold", fontSize=16, leading=19)

    @staticmethod
    def _docx_heading(doc: Document, value: str) -> None: doc.add_heading(value, level=1)
    @staticmethod
    def _dates(start: str | None, end: str | None) -> str: return " – ".join(value for value in [start, end] if value)
    @staticmethod
    def _contact(identity: ApplicationIdentitySnapshot) -> str: return " | ".join(value for value in [identity.email, identity.phone, identity.location, identity.linkedin_url, identity.github_url, identity.portfolio_url] if value)


def _paragraph(value: str) -> str:
    return re.sub(r"\n+", "<br/>", value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
